"""Simple BEAM_BRACE_TIE input/planning path; no assembly pose inference."""
import hashlib
import json
import math
from pathlib import Path
from openpyxl import load_workbook
from .analysis_defaults import open_analysis_config, CONFIG_ID, Mapping, deepcopy

from .analysis_project import validate_project_files, _components, _one
from .analysis_plan import semantic_name
from .main_frame_assembly import _read_named_points
from .runtime_compatibility import assert_runtime_compatible
from .column_column_tie import build_column_plan
from .column_hoop_tie import build_hoop_plan
from .column_beam_coupling import build_coupling_plan

FIELDS = ('rule_id', 'enabled', 'beam_half_length_mm', 'brace_end_length_mm', 'beam_shell_side', 'brace_shell_side')
ROLE_STATIONS = (('BRACE_FRONT', 'C', 'B'), ('BRACE_REAR', 'E', 'D'))


def read_tie_config(path):
    wb = open_analysis_config(path, load_workbook, 'Tie_Config', 'BEAM_BRACE_TIE', FIELDS, data_only=False, read_only=True)
    try:
        if 'Tie_Config' not in wb.sheetnames:
            raise ValueError('Missing Tie_Config sheet')
        rows = list(wb['Tie_Config'].iter_rows(values_only=True))
        if not rows or tuple(rows[0]) != FIELDS:
            raise ValueError('Tie_Config headers must be: '+', '.join(FIELDS))
        records = [r for r in rows[1:] if any(v is not None for v in r)]
        if len(records) != 1:
            raise ValueError('Tie_Config requires exactly one BEAM_BRACE_TIE row')
        data = dict(zip(FIELDS, records[0]))
        if data['rule_id'] != 'BEAM_BRACE_TIE':
            raise ValueError('Unsupported Tie rule')
        enabled = data['enabled']
        if enabled not in (True, False, 'TRUE', 'FALSE'):
            raise ValueError('enabled must be TRUE or FALSE')
        data['enabled'] = enabled in (True, 'TRUE')
        for key in ('beam_half_length_mm', 'brace_end_length_mm'):
            v = data[key]
            if isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) or v <= 0:
                raise ValueError('Expected positive numeric '+key)
            data[key] = float(v)
        for key in ('beam_shell_side', 'brace_shell_side'):
            data[key] = data[key] or 'SIDE2'
            if data[key] not in ('SIDE1', 'SIDE2'):
                raise ValueError('Invalid '+key)
        return data
    finally:
        wb.close()


def local_stations(path):
    """Read existing GC_mm / GE_mm directly; do not infer any Assembly pose."""
    wb = load_workbook(path, data_only=True, read_only=True)
    found = {}
    try:
        for sheet in wb:
            cols = None
            for row in sheet.iter_rows(values_only=True):
                if all(h in row for h in ('参数名', '数值', '单位')):
                    cols = {h: row.index(h) for h in ('参数名', '数值', '单位')}
                    continue
                if cols is None:
                    continue
                key = row[cols['参数名']]
                if key not in ('GC_mm', 'GE_mm'):
                    continue
                v, units = row[cols['数值']], row[cols['单位']]
                if isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) or units != 'mm':
                    raise ValueError('Expected cached numeric '+key+' in mm')
                if key in found and abs(found[key]-v/1000.) > 1.e-9:
                    raise ValueError('Conflicting coordinate parameter '+key)
                found[key] = v/1000.
    finally:
        wb.close()
    if set(found) != {'GC_mm', 'GE_mm'}:
        raise ValueError('Missing real GC_mm / GE_mm; explicit local stations required')
    return {'C': found['GC_mm'], 'E': found['GE_mm']}


def build_plan(files, analysis_path):
    summary = validate_project_files(files)
    config = read_tie_config(analysis_path)
    plan = dict(project_id=summary['project_id'], model_name=summary['model_name'], rule_id='BEAM_BRACE_TIE',
                config=config, instances=[], connections=[], partitions={}, source_files={k: str(p) for k,p in files.paths().items()})
    if isinstance(analysis_path,Mapping):
        plan['source_files']['analysis_config'] = CONFIG_ID
        plan['analysis_config'] = deepcopy(dict(analysis_path))
    else:
        plan['source_files']['analysis'] = str(Path(analysis_path).resolve())
    if not config['enabled']:
        return plan
    selected = []
    for role in ('INCLINED_BEAM', 'BRACE_FRONT', 'BRACE_REAR'):
        matches = [m for m in summary['instances'] if (m.get('canonical_role') or m.get('component_code')) == role]
        if len(matches) != 1:
            raise ValueError('Expected unique canonical role '+role)
        member = dict(matches[0], role=role)
        users = [m['instance_name'] for m in summary['instances'] if m['part_name'] == member['part_name']]
        if len(users) != 1:
            raise ValueError('Shared Part would affect other instances: '+member['part_name'])
        selected.append(member)
    geometry = _components(files.components, selected, 1.e-6)
    for m in selected:
        m['component_geometry'] = geometry[m['part_name']]
    stations = local_stations(files.coordinate)
    points, _ = _read_named_points(files.coordinate)
    for name in ('B', 'C', 'D', 'E'):
        expected = summary.get('assembly_points', {}).get(name)
        actual = points.get(name)
        if expected is None or actual is None or len(expected) != 3:
            raise ValueError('Missing real connection point '+name)
        if any(not isinstance(v,(int,float)) or not math.isfinite(v) for v in expected+list(actual)):
            raise ValueError('Invalid connection point '+name)
        if max(abs(a-b) for a,b in zip(expected,actual)) > 1.e-6:
            raise ValueError('Coordinate/summary point mismatch: '+name)
    half = config['beam_half_length_mm']/1000.
    end = config['brace_end_length_mm']/1000.
    beam = selected[0]
    beam_length = beam['component_geometry']['length_m']
    ranges = []
    plan['partitions'][beam['part_name']] = []
    for brace, (role, station, start) in zip(selected[1:], ROLE_STATIONS):
        length = brace['component_geometry']['length_m']
        low, high = stations[station]-half, stations[station]+half
        if not 0 < low < high < beam_length or not 0 < end < length/2.:
            raise ValueError('Partition range outside Part or overlapping brace end zones')
        if any(max(low,a) < min(high,b) for a,b in ranges):
            raise ValueError('Beam connection patches overlap')
        ranges.append((low, high))
        plan['partitions'][beam['part_name']].extend([low, high])
        plan['partitions'][brace['part_name']] = [end, length-end]
        suffix = role.replace('BRACE_', '')
        plan['connections'].append(dict(id='BEAM_BRACE_'+suffix, station=station, station_m=stations[station],
            brace_start=start, brace_end=station, beam=beam['instance_name'], brace=brace['instance_name'],
            beam_part=beam['part_name'], brace_part=brace['part_name'], beam_range=[low,high],
            brace_length=length, brace_range=[length-end,length],
            beam_region=semantic_name('REGION_BEAM_BRACE_'+suffix),
            brace_region=semantic_name('REGION_BRACE_'+suffix+'_BEAM'),
            tie=semantic_name('TIE_BEAM_BRACE_'+suffix), master=brace['instance_name'], slave=beam['instance_name']))
    # Runtime needs names, dimensions and ranges, never placement transforms.
    plan['instances'] = [{k:m[k] for k in ('instance_name','part_name','role','component_geometry')} for m in selected]
    return plan


def generate(files, analysis_path=None, output_dir=None, final_output_dir=None):
    if analysis_path is None:
        folder = files.assembly_summary.parent
        analysis_path = _one([p for p in folder.iterdir() if p.is_file() and not p.name.startswith('~$')
                            and (p.name == 'analysis.xlsx' or p.name.endswith('_analysis.xlsx'))], 'analysis.xlsx')
    plan = build_plan(files, analysis_path)
    if final_output_dir is not None and not isinstance(analysis_path,Mapping):
        plan['source_files']['analysis'] = str(Path(final_output_dir).resolve()/(plan['project_id']+'_analysis.xlsx'))
    encoded = json.dumps(plan, sort_keys=True, ensure_ascii=True, allow_nan=False)
    fingerprint = hashlib.sha256(encoded.encode('utf-8')).hexdigest()
    column = build_column_plan(files, analysis_path)
    if column is not None:
        plan['column_column_tie'] = column
    hoop = build_hoop_plan(files, analysis_path)
    if hoop is not None:
        plan['column_hoop_tie'] = hoop
    coupling = build_coupling_plan(files, analysis_path)
    if coupling is not None:
        plan['column_beam_coupling'] = coupling
    encoded = json.dumps(plan, sort_keys=True, ensure_ascii=True, allow_nan=False)
    output = Path(output_dir) if output_dir else files.assembly_summary.parent/'step05_validation'
    output.mkdir(parents=True, exist_ok=True)
    project = plan['project_id']
    if any(c in project for c in '/\\<>:"|?*'):
        raise ValueError('Unsafe project id')
    paths = dict(script=output/(project+'_analysis_setup.py'), plan=output/(project+'_analysis_plan.json'),
                 runtime_report=output/(project+'_analysis_setup_report.json'))
    if final_output_dir is not None:
        paths['runtime_report'] = Path(final_output_dir)/(project+'_analysis_setup_report.json')
    source = '# -*- coding: utf-8 -*-\nimport json\nPLAN = json.loads(%r)\n' % encoded
    source += 'FINGERPRINT = %r\nREPORT = %r\n' % (fingerprint, str(paths['runtime_report'].resolve()).replace('\\','/'))
    beam_source = Path(__file__).with_name('beam_brace_tie_runtime.py').read_text(encoding='utf-8')
    if column is None and hoop is None and coupling is None:
        source += Path(__file__).with_name('tie_preflight_runtime.py').read_text(encoding='utf-8') + '\n'
        source += beam_source.replace('    execute(PLAN,mdb,REPORT,FINGERPRINT,XYPLANE,ON,COMPUTED)',
            '    tie_batch = Step05TieBatch()\n    execute(PLAN,mdb,REPORT,FINGERPRINT,XYPLANE,ON,COMPUTED,tie_batch)\n    tie_batch.finish(REPORT)')
    else:
        # Preserve every beam function verbatim; extend only the generated entry point.
        body, entry = beam_source.split("\nif __name__ == '__main__':", 1)
        source += Path(__file__).with_name('tie_preflight_runtime.py').read_text(encoding='utf-8') + '\n'
        source += body + '\n' + Path(__file__).with_name('column_axis.py').read_text(encoding='utf-8')
        source += '\n' + Path(__file__).with_name('column_column_tie_runtime.py').read_text(encoding='utf-8')
        if column is not None:
            column_fingerprint = hashlib.sha256(json.dumps(column, sort_keys=True).encode('utf-8')).hexdigest()
            source += '\nCOLUMN_FINGERPRINT = %r\n' % column_fingerprint
        if hoop is not None:
            source += '\n' + Path(__file__).with_name('column_hoop_tie_runtime.py').read_text(encoding='utf-8')
            hoop_fingerprint = hashlib.sha256(json.dumps(hoop, sort_keys=True).encode('utf-8')).hexdigest()
            source += '\nHOOP_FINGERPRINT = %r\n' % hoop_fingerprint
        if coupling is not None:
            source += '\n' + Path(__file__).with_name('column_beam_coupling_runtime.py').read_text(encoding='utf-8')
        source += "\nif __name__ == '__main__':"
        execution = (
            "    tie_batch = Step05TieBatch()\n"
            "    previous_rules = {}\n"
            "    if os.path.isfile(REPORT):\n"
            "        with open(REPORT) as handle:\n"
            "            previous_rules = json.load(handle)\n"
            "    execute(PLAN,mdb,REPORT,FINGERPRINT,XYPLANE,ON,COMPUTED,tie_batch)\n")
        if coupling is not None:
            execution = execution.replace('    execute(PLAN', '    from abaqusConstants import OFF, DISTRIBUTING, WHOLE_SURFACE, UNIFORM\n    coupling_context = prepare_coupling_geometry(PLAN,mdb,REPORT,XYPLANE,ON)\n    execute(PLAN')
        if hoop is not None:
            execution += "    hoop_context = prepare_hoop(PLAN,mdb,REPORT,HOOP_FINGERPRINT,previous_rules.get('column_hoop_tie'),XYPLANE,ON)\n"
        if column is not None:
            execution += "    execute_column(PLAN,mdb,REPORT,COLUMN_FINGERPRINT,previous_rules.get('column_column_tie'),XYPLANE,ON,COMPUTED,tie_batch)\n"
        if hoop is not None:
            execution += "    execute_hoop(hoop_context,ON,COMPUTED,tie_batch)\n"
        execution += '    tie_batch.finish(REPORT)\n'
        if coupling is not None:
            execution += '    execute_coupling(coupling_context,ON,OFF,DISTRIBUTING,WHOLE_SURFACE,UNIFORM)\n'
        source += entry.replace('    execute(PLAN,mdb,REPORT,FINGERPRINT,XYPLANE,ON,COMPUTED)',execution.rstrip())
    assert_runtime_compatible(source, str(paths['script']))
    compile(source, str(paths['script']), 'exec')
    paths['script'].write_text(source, encoding='utf-8')
    paths['plan'].write_text(json.dumps(plan, indent=2, ensure_ascii=False), encoding='utf-8')
    return paths
