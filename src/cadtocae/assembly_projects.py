"""Step04 folder discovery and batch orchestration; no assembly placement rules."""
from __future__ import annotations

import copy
import json
import shutil
import traceback
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook

from .assembly_script import generate_assembly_scripts_from_workbook
from .main_frame_assembly import build_payload, detect_structure_type, read_components_payload, read_step02_model_name

SUFFIXES = {"part_script": "_create_parts_in_cae.py", "coordinate": "_coordinate.xlsx"}


def _candidates(folder: Path, recursive: bool) -> dict:
    projects: dict = {}
    for path in sorted(folder.rglob("*") if recursive else folder.iterdir()):
        if not path.is_file() or path.name.startswith("~$"):
            continue
        for role, suffix in SUFFIXES.items():
            if path.name.lower().endswith(suffix):
                project = path.name[:-len(suffix)]
                if project:
                    projects.setdefault(project, {}).setdefault(role, []).append(str(path.resolve()))
    return projects


def _identity_check(metadata: dict, project: str, structure: str) -> None:
    for key in ("project_id", "project_code", "project_prefix", "model_name"):
        value = metadata.get(key)
        if value not in (None, "") and str(value).strip() != project:
            raise ValueError("Project ID Mismatch: %s=%s; filename=%s" % (key, value, project))
    value = metadata.get("structure_type")
    if value not in (None, "") and str(value).strip().upper() != structure:
        raise ValueError("Project ID Mismatch: structure_type=%s; filename structure=%s" % (value, structure))


def _coordinate_identity(path: Path, project: str, structure: str) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        # Legacy coordinate workbooks may have no explicit project ID. Check all
        # available identity fields, plus the primary sheet's structural title.
        title = str(workbook.worksheets[0].cell(1, 1).value or "")
        labels = {"单桩单立柱": "SP_SC", "单桩双立柱": "SP_DC", "双桩双立柱": "DP"}
        for label, expected in labels.items():
            if label in title and expected != structure:
                raise ValueError("Project ID Mismatch: coordinate title identifies %s, expected %s" % (expected, structure))
        keys = {"project_id", "project_code", "project_prefix", "model_name", "structure_type"}
        for sheet in workbook:
            for row in sheet.iter_rows(values_only=True):
                if row and isinstance(row[0], str) and row[0].strip().lower() in keys:
                    # Metadata rows use key/value or parameter/description/value.
                    value = row[2] if len(row) > 2 and row[2] not in (None, "") else (row[1] if len(row) > 1 else None)
                    _identity_check({row[0].strip().lower(): value}, project, structure)
    finally:
        workbook.close()


def inspect_project(project: str, candidates: dict) -> dict:
    row = dict(project_id=project, structure_type="Unknown", candidates=candidates,
               input_status="Ready", capability="Supported", generation_status="Ready",
               result="", warnings=[], errors=[], input_files={}, planned_instances=0)
    stage = "Invalid Part Metadata"
    try:
        try:
            row["structure_type"] = detect_structure_type(project)
        except ValueError as exc:
            row.update(capability="Unsupported", input_status="Unknown Structure")
            raise ValueError(str(exc))
        if any(len(paths) > 1 for paths in candidates.values()):
            row["input_status"] = "Ambiguous Files"
            raise ValueError("Multiple candidates; no file chosen:\n" + "\n".join(p for paths in candidates.values() for p in paths))
        for role, label in (("part_script", "Missing Part Script"), ("coordinate", "Missing Coordinate")):
            if not candidates.get(role):
                row["input_status"] = label
                raise ValueError(label + ": " + project + SUFFIXES[role])
        row["input_files"] = {key: paths[0] for key, paths in candidates.items()}
        part, coordinate = (Path(row["input_files"][key]) for key in ("part_script", "coordinate"))
        if part.parent != coordinate.parent:
            row["input_status"] = "Ambiguous Files"
            raise ValueError("Inputs are in different folders; put the matching pair in one project folder.\n%s\n%s" % (part, coordinate))
        row["source_folder"] = str(coordinate.parent)
        data = read_components_payload(part)  # Existing embedded COMPONENTS_JSON parser; never execute Python.
        model = read_step02_model_name(part)
        if not model:
            raise ValueError("Step02 MODEL_NAME is missing.")
        _identity_check({"model_name": model}, project, row["structure_type"])
        for metadata in (data, data.get("meta", {}), data.get("metadata", {})):
            _identity_check(metadata, project, row["structure_type"])
        components = data.get("components")
        if not isinstance(components, list) or not components:
            raise ValueError("Step02 components are missing or empty.")
        stage = "Invalid Coordinate"
        _coordinate_identity(coordinate, project, row["structure_type"])
        stage = "Preflight Error"
        payload = build_payload(coordinate, part, project_code=project, model_name=project)
        if payload.get("errors"):
            raise ValueError("; ".join(payload["errors"]))
        row["warnings"] = list(payload.get("warnings", []))
        row["planned_instances"] = len(payload["instance_plan"])
        row["result"] = "%d instances planned%s" % (row["planned_instances"], "; %d warnings" % len(row["warnings"]) if row["warnings"] else "")
    except Exception as exc:
        if row["input_status"] == "Ready":
            row["input_status"] = "Project ID Mismatch" if "Mismatch" in str(exc) or "match" in str(exc).lower() else stage
        if row["capability"] != "Unsupported":
            row["capability"] = "Needs Review"
        row.update(generation_status="Failed", result=str(exc).splitlines()[0][:160], errors=[str(exc)])
    return row


def scan_projects(folder: str | Path, recursive: bool = False) -> dict:
    root = Path(folder).resolve()
    if not root.is_dir():
        raise ValueError("Project Folder is not a directory: %s" % root)
    return {"folder": str(root), "recursive": recursive,
            "projects": [inspect_project(project, candidates) for project, candidates in sorted(_candidates(root, recursive).items())]}


def can_generate(row: dict) -> bool:
    return row["input_status"] == "Ready" and row["capability"] == "Supported"


def project_details(row: dict) -> str:
    lines = [row["project_id"], "%s / %s / %s" % (row["input_status"], row["capability"], row["generation_status"])]
    for role, paths in row["candidates"].items():
        lines.extend(role + ": " + path for path in paths)
    lines.extend("WARNING: " + value for value in row["warnings"])
    lines.extend("ERROR: " + value for value in row["errors"])
    lines.extend("Output: " + value for value in row.get("generated_files", []))
    if row.get("report_path"):
        lines.append("Debug Report: " + row["report_path"])
    return "\n".join(lines)


def generate_batch(scan: dict, selected: set[str] | None = None, progress: Callable | None = None) -> dict:
    result = copy.deepcopy(scan)
    counts = dict(total=0, success=0, warning=0, failed=0, skipped=0)
    targets = [r for r in result["projects"] if selected is None or r["project_id"] in selected]
    counts["total"] = len(targets)

    def emit(event, row, index):
        if progress:
            progress({"event": event, "row": copy.deepcopy(row), "index": index, "total": len(targets)})

    for index, row in enumerate(targets, 1):
        if selected is None and not can_generate(row):
            counts["skipped"] += 1
            emit("skipped", row, index)
            continue
        row["generation_status"] = "Queued"
        emit("queued", row, index)
    for index, row in enumerate(targets, 1):
        if selected is None and row["generation_status"] != "Queued":
            continue
        row["generation_status"] = "Generating"
        row.pop("generated_files", None)
        row.pop("report_path", None)
        emit("started", row, index)
        try:
            # Rediscover candidate paths as well as metadata: added duplicates,
            # renamed files and edits after Scan must all be caught.
            candidates = _candidates(Path(scan["folder"]), scan["recursive"]).get(row["project_id"], {})
            fresh = inspect_project(row["project_id"], candidates)
            row.update(fresh)
            if not can_generate(row):
                raise ValueError("; ".join(row["errors"]))
            row["generation_status"] = "Generating"
            emit("preflight", row, index)
            output = generate_assembly_scripts_from_workbook(
                row["input_files"]["coordinate"], row["source_folder"],
                components_json=row["input_files"]["part_script"], overwrite=True,
            )
            # Keep the direct-file API and its debug directory intact. The folder
            # workflow also publishes the requested report beside project inputs.
            report = Path(row["source_folder"]) / (row["project_id"] + "_step04_assembly_script_report.json")
            shutil.copy2(output.report_path, report)
            row["report_path"] = str(report)
            if output.status == "failed":
                raise ValueError("; ".join(output.messages))
            row["warnings"] = json.loads(report.read_text(encoding="utf-8")).get("warnings", [])
            row["generated_files"] = [*output.script_paths, output.summary_path]
            row["generation_status"] = "Generated"
            row["result"] = "Generated — %d instances%s" % (row["planned_instances"], "; %d warnings" % len(row["warnings"]) if row["warnings"] else "")
            counts["warning" if row["warnings"] else "success"] += 1
        except Exception as exc:
            row.update(generation_status="Failed", result=str(exc).splitlines()[0][:160], errors=[str(exc)])
            # Ambiguous projects have no unique output owner; keep their workflow
            # error report in the selected folder and list every candidate.
            report = Path(row.get("source_folder", scan["folder"])) / (row["project_id"] + "_step04_assembly_script_report.json")
            try:
                report.write_text(json.dumps({"status": "failed", "project_id": row["project_id"],
                    "candidates": row["candidates"], "errors": row["errors"], "traceback": traceback.format_exc()}, ensure_ascii=False, indent=2), encoding="utf-8")
                row["report_path"] = str(report)
            except OSError as report_error:
                row["errors"].append("Cannot write report %s: %s" % (report, report_error))
            counts["failed"] += 1
        emit("finished", row, index)
    result["generation_counts"] = counts
    return result
