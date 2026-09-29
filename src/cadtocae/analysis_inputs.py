"""Strict, side-effect-free Step05 A1 input readers (JSON or two-sheet Excel)."""
from __future__ import annotations

import json
import math
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

REGION_COLUMNS = ("region_id", "connection_id", "target_role", "reference_station",
                  "half_length_m", "partition_rule", "enabled")


def _text(value: Any, field: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError("Missing/text required: " + field)
    return value.strip()


def _enabled(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (str, int)) and str(value).strip().upper() in ("TRUE", "ON", "1"):
        return True
    if isinstance(value, (str, int)) and str(value).strip().upper() in ("FALSE", "OFF", "0"):
        return False
    raise ValueError("enabled must be TRUE/FALSE or ON/OFF")


def normalize_inputs(data: dict[str, Any]) -> dict[str, Any]:
    info = dict(data.get("Project_Info") or {})
    project = _text(info.get("project_id"), "project_id")
    if str(info.get("schema_version", "")) != "1":
        raise ValueError("Unsupported analysis schema_version; expected 1")
    if info.get("length_unit") != "m":
        raise ValueError("A1/A2 length_unit must explicitly be m")
    try:
        tolerance = float(info.get("geometry_tolerance_m", 1.e-6))
    except (ValueError, TypeError):
        raise ValueError("Invalid geometry_tolerance_m") from None
    if not math.isfinite(tolerance) or tolerance <= 0 or tolerance > 1.e-4:
        raise ValueError("geometry_tolerance_m must be in (0, 0.0001]")
    rows, ids, connections = [], set(), set()
    for index, raw in enumerate(data.get("Partition_Region", []), 2):
        if not isinstance(raw, dict):
            raise ValueError("Partition_Region row must be an object")
        row = {key: _text(raw.get(key), key) for key in REGION_COLUMNS
               if key not in ("enabled", "half_length_m")}
        row["enabled"] = _enabled(raw.get("enabled"))
        try:
            half = float(raw["half_length_m"])
        except (ValueError, TypeError, KeyError):
            raise ValueError("Invalid half_length_m") from None
        if isinstance(raw["half_length_m"], bool) or not math.isfinite(half) or half <= tolerance:
            raise ValueError("half_length_m must be finite and greater than geometry tolerance")
        row["half_length_m"] = half
        for optional in ("instance_id", "connection_group", "brace_instance_id"):
            if raw.get(optional) not in (None, ""):
                row[optional] = _text(raw[optional], optional)
        if row["region_id"] in ids or row["connection_id"] in connections:
            raise ValueError("Duplicate region or connection definition")
        ids.add(row["region_id"])
        connections.add(row["connection_id"])
        row["source_input"] = "Partition_Region:%d" % index
        rows.append(row)
    if not any(row["enabled"] for row in rows):
        raise ValueError("No enabled partition definition")
    return {"Project_Info": {"project_id": project, "schema_version": 1,
                             "length_unit": "m", "geometry_tolerance_m": tolerance},
            "Partition_Region": rows}


def read_analysis_inputs(path: str | Path) -> dict[str, Any]:
    path = Path(path)
    if path.suffix.lower() == ".json":
        return normalize_inputs(json.loads(path.read_text(encoding="utf-8-sig")))
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Expected analysis .json or .xlsx")
    # Own the file handle as well as the workbook, including exceptional exits.
    with path.open("rb") as stream:
        workbook = load_workbook(stream, read_only=True, data_only=False)
        try:
            if set(workbook.sheetnames) != {"Project_Info", "Partition_Region"}:
                raise ValueError("A1/A2 requires exactly Project_Info and Partition_Region")
            tables = {}
            for sheet in workbook:
                values = list(sheet.values)
                if any(isinstance(cell, str) and cell.startswith("=") for row in values for cell in row):
                    raise ValueError("Analysis input formulas are not supported; supply literal values")
                tables[sheet.title] = values
            info = {}
            if tuple(tables["Project_Info"][0][:2]) != ("key", "value"):
                raise ValueError("Project_Info header must be key,value")
            for row in tables["Project_Info"][1:]:
                if not any(value is not None for value in row):
                    continue
                key = _text(row[0], "Project_Info.key")
                if key in info:
                    raise ValueError("Duplicate Project_Info key: " + key)
                info[key] = row[1]
            headers = tables["Partition_Region"][0]
            if len(set(headers)) != len(headers) or not set(REGION_COLUMNS).issubset(headers):
                raise ValueError("Missing or duplicate Partition_Region columns")
            rows = [dict(zip(headers, row)) for row in tables["Partition_Region"][1:]
                    if any(value is not None for value in row)]
            return normalize_inputs({"Project_Info": info, "Partition_Region": rows})
        finally:
            workbook.close()
