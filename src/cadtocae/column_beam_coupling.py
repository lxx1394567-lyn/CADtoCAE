"""Independent RP/two-coupling planning; verified Tie rules remain frozen."""
import json
import math
import tempfile
from pathlib import Path
from openpyxl import load_workbook
from .analysis_defaults import open_analysis_config
from .analysis_project import validate_project_files
from .main_frame_assembly import _read_named_points
from .column_axis import cct_placement_axis
from .workbook import export_abaqus_json

FIELDS = ('rule_id','enabled','column_top_length_mm','coupling_type','u1','u2','u3','ur1','ur2','ur3')


def read_coupling_config(path):
    wb = open_analysis_config(path, load_workbook, 'Coupling_Config', 'COLUMN_BEAM_COUPLING', FIELDS, read_only=True, data_only=False)
    try:
        if 'Coupling_Config' not in wb.sheetnames:
            return None
        rows = list(wb['Coupling_Config'].values)
        if len(rows) != 2 or rows[0] != FIELDS:
            raise ValueError('Coupling_Config requires '+', '.join(FIELDS))
        config = dict(zip(FIELDS,rows[1]))
        if config['rule_id'] != 'COLUMN_BEAM_COUPLING':
            raise ValueError('Unknown coupling rule')
        for key in ('enabled','u1','u2','u3','ur1','ur2','ur3'):
            value=config[key]
            if type(value) not in (bool,str) or value not in (True,False,'TRUE','FALSE'):
                raise ValueError('Expected TRUE/FALSE for '+key)
            config[key]=value in (True,'TRUE')
        if not config['enabled']:
            return None
        if config['coupling_type'] != 'DISTRIBUTING' or not all(config[k] for k in FIELDS[4:]):
            raise ValueError('First coupling rule requires DISTRIBUTING and all six DOFs ON')
        length=config['column_top_length_mm']
        if length is None: length=50.
        if isinstance(length,bool) or not isinstance(length,(int,float)) or not math.isfinite(length) or length<=0:
            raise ValueError('Invalid column_top_length_mm')
        config['column_top_length_mm']=float(length)
        return config
    finally:
        wb.close()


def build_coupling_plan(files, path):
    config=read_coupling_config(path)
    if config is None: return None
    summary=validate_project_files(files)
    with tempfile.TemporaryDirectory(prefix='step05_coupling_') as tmp:
        exported=export_abaqus_json(files.components,Path(tmp)/'components.json',selection='complete')
        components=json.loads(exported.read_text(encoding='utf-8'))['components']
    points,_=_read_named_points(files.coordinate)
    f=summary.get('assembly_points',{}).get('F',points.get('F'))
    if f is None or len(f)!=3 or not all(isinstance(x,(int,float)) and math.isfinite(x) for x in f):
        raise ValueError('Missing/invalid assembly point F')
    if 'F' in points and max(abs(a-b) for a,b in zip(f,points['F']))>1.e-6:
        raise ValueError('Coordinate/summary point F mismatch')
    wb=load_workbook(files.coordinate,read_only=True,data_only=True)
    stations=[]
    try:
        for sheet in wb:
            rows=iter(sheet.values)
            for row in rows:
                if '参数名' in row and '数值' in row and '单位' in row:
                    names={v:i for i,v in enumerate(row)}
                    for data in rows:
                        if data[names['参数名']]=='GF_mm':
                            value=data[names['数值']]
                            if data[names['单位']]!='mm' or isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value):
                                raise ValueError('Invalid cached GF_mm')
                            stations.append(value/1000.)
                    break
    finally: wb.close()
    if not stations or max(stations)-min(stations)>1.e-9: raise ValueError('Missing/conflicting GF_mm')
    return resolve_coupling(summary,components,f,stations[0],config)


def resolve_coupling(summary,components,f,station,config):
    members=[]
    for role,kind in (('COLUMN_UP','PIPE'),('INCLINED_BEAM','C_CHANNEL')):
        candidates=[m for m in summary['instances'] if (m.get('canonical_role') or m.get('component_code'))==role]
        if len(candidates)!=1: raise ValueError('Expected unique '+role)
        m=candidates[0]
        if len([x for x in summary['instances'] if x['part_name']==m['part_name']])!=1:
            raise ValueError('Coupling requires independent Part ownership')
        matches=[c for c in components if c['part_name']==m['part_name'] and c['component_code']==m['component_code']]
        if len(matches)!=1: raise ValueError('Missing/ambiguous coupling component')
        g=matches[0]
        if g['section_kind']!=kind or g['model_policy']!='SHELL': raise ValueError('Unsupported coupling geometry '+role)
        length=g['length_m']
        if not math.isfinite(length) or length<=0 or abs(m['geometry_info']['length']-length)>1.e-6:
            raise ValueError('Coupling component length mismatch')
        placement=dict(m['placement'])
        if m.get('axis_direction') is not None: placement.setdefault('axis_direction',m['axis_direction'])
        axis=list(cct_placement_axis(placement))
        member=dict(role=role,instance_name=m['instance_name'],part_name=m['part_name'],length_m=length,
            axis_direction=axis,placement=placement,section_params_m=g['section_params_m'])
        if role=='COLUMN_UP':
            origin=placement.get('final_origin',placement['translation'])
            if len(origin)!=3 or not all(math.isfinite(x) for x in origin): raise ValueError('Invalid column origin')
            radius=(g['section_params_m']['od_m']-g['thickness_m'])/2.
            if radius<=0: raise ValueError('Invalid column radius')
            member.update(origin=list(origin),radius_m=radius,thickness_m=g['thickness_m'])
        members.append(member)
    up,beam=members
    if not 0<station<beam['length_m']: raise ValueError('F station outside beam')
    rp=[up['origin'][i]+up['axis_direction'][i]*up['length_m'] for i in range(3)]
    return dict(rule_id='COLUMN_BEAM_COUPLING',config=config,column=up,beam=beam,
        rp_name='STEP05_RP_COLUMN_UP',rp_coordinate=rp,f_point=list(f),f_station_m=station,
        source_set='SET_BEAM_SEC_F',target_type='EDGE',column_top_length_used=False,
        targets=[dict(instance_name=up['instance_name'],region='STEP05_REGION_COLUMN_UP_TOP',surface='STEP05_SURF_COLUMN_UP_TOP',coupling='STEP05_COUPLING_COLUMN_UP'),
                 dict(instance_name=beam['instance_name'],region='STEP05_REGION_BEAM_F',surface='STEP05_SURF_BEAM_F',coupling='STEP05_COUPLING_BEAM_F')])
