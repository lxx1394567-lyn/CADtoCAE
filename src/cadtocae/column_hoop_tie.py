"""One canonical hoop pair; all contact dimensions come from normalized inputs."""
import json
import math
import tempfile
from pathlib import Path
from openpyxl import load_workbook
from .analysis_defaults import open_analysis_config
from .analysis_project import validate_project_files
from .workbook import export_abaqus_json
from .column_axis import cct_placement_axis, cct_dot
from .main_frame_assembly import is_hoop_component_code
from .analysis_capability import is_hoop_member, match_hoop_component, hoop_mapping


def hoop_enabled(path):
    wb = open_analysis_config(path,load_workbook,'Column_Hoop_Tie_Config','COLUMN_HOOP_TIE',('rule_id','enabled'),read_only=True,data_only=False)
    try:
        if 'Column_Hoop_Tie_Config' not in wb.sheetnames:
            return False
        rows = list(wb['Column_Hoop_Tie_Config'].values)
        if len(rows) != 2 or rows[0] != ('rule_id','enabled') or rows[1][0] != 'COLUMN_HOOP_TIE':
            raise ValueError('Column_Hoop_Tie_Config requires rule_id, enabled and COLUMN_HOOP_TIE')
        value = rows[1][1]
        if type(value) not in (bool,str) or value not in (True,False,'TRUE','FALSE'):
            raise ValueError('Column_Hoop_Tie_Config enabled must be TRUE/FALSE')
        return value in (True,'TRUE')
    finally:
        wb.close()


def build_hoop_plan(files,path):
    if not hoop_enabled(path):
        return None
    summary = validate_project_files(files)
    with tempfile.TemporaryDirectory(prefix='step05_hoop_') as tmp:
        source = export_abaqus_json(files.components,Path(tmp)/'components.json',selection='complete')
        components = json.loads(source.read_text(encoding='utf-8'))['components']
    return resolve_hoop(summary,components)


def resolve_hoop(summary,components):
    def one(values,label):
        if len(values) != 1:
            raise ValueError('Expected unique '+label)
        return values[0]
    def geometry(member,kind,policy):
        if kind == 'HOOP_BAND':
            g = match_hoop_component(member,components)
        else:
            g = one([c for c in components if c['part_name'] == member['part_name'] and c['component_code'] == member['component_code']], 'normalized component')
        if g['section_kind'] != kind or g['model_policy'] != policy:
            raise ValueError('Expected '+kind+' '+policy)
        return g
    column = one([m for m in summary['instances'] if (m.get('canonical_role') or m.get('component_code')) == 'COLUMN_DOWN'],'COLUMN_DOWN')
    one([m for m in summary['instances'] if m['part_name'] == column['part_name']],'column Part owner')
    g = geometry(column,'PIPE','SHELL')
    params = g['section_params_m']
    axis = cct_placement_axis(column['placement'])
    origin = column['placement'].get('final_origin',column['placement']['translation'])
    length = g['length_m']
    if abs(length-column['geometry_info']['length']) > 1.e-6:
        raise ValueError('Column workbook/summary length mismatch')
    member = dict(role='COLUMN_DOWN',part_name=column['part_name'],instance_name=column['instance_name'],
        length_m=length,radius_m=(params['od_m']-g['thickness_m'])/2.,thickness_m=g['thickness_m'],
        origin=origin,axis_direction=list(axis),physical_side='OUTER')
    hoops = [m for m in summary['instances'] if is_hoop_member(m)]
    groups = {m.get('connection_group') for m in hoops}
    if len(hoops) != 2 or len(groups) != 1 or None in groups or '' in groups:
        raise ValueError('COLUMN_HOOP_TIE requires exactly one complete HOOP group with two instances')
    halves = []
    for label in ('A','B'):
        h = one([h for h in hoops if (h.get('pair_member') or h['instance_name'].rsplit('_',1)[-1]) == label],'HOOP pair member '+label)
        hg = geometry(h,'HOOP_BAND','SOLID')
        hp = hg['section_params_m']
        haxis = cct_placement_axis(h['placement'])
        if abs(abs(cct_dot(haxis,axis))-1.) > 1.e-10:
            raise ValueError('HOOP extrusion axis is not parallel to column')
        start = cct_dot(h['placement']['translation'],axis)
        end = start+hp['width_m']*cct_dot(haxis,axis)
        halves.append(dict(label=label,instance_name=h['instance_name'],part_name=h['part_name'],
            component_mapping=hoop_mapping(h,hg),
            connection_group=h['connection_group'],placement=h['placement'],width_m=hp['width_m'],
            master=member['instance_name'],slave=h['instance_name'],master_physical_side='OUTER',slave_physical_side='INNER',
            inner_radius_m=hp['inner_radius_m'],thickness_m=hg['thickness_m'],expected_band=sorted([start,end]),
            column_region='STEP05_REGION_COLUMN_DOWN_HOOP_'+label,
            hoop_region='STEP05_REGION_HOOP_'+label+'_COLUMN_DOWN',tie='STEP05_TIE_COLUMN_DOWN_HOOP_'+label))
    for value in (length,member['radius_m'],member['thickness_m']):
        if not math.isfinite(value) or value <= 0: raise ValueError('Invalid PIPE dimensions')
    for half in halves:
        for key in ('width_m','inner_radius_m','thickness_m'):
            if not math.isfinite(half[key]) or half[key] <= 0: raise ValueError('Invalid HOOP dimensions')
    if max(abs(a-b) for a,b in zip(halves[0]['expected_band'],halves[1]['expected_band'])) > 1.e-6:
        raise ValueError('HOOP A/B axial bands differ')
    if abs(halves[0]['inner_radius_m']-halves[1]['inner_radius_m']) > 1.e-6:
        raise ValueError('HOOP A/B inner radii differ')
    start = cct_dot(origin,axis)
    band = halves[0]['expected_band']
    overlap = [max(start,band[0]),min(start+length,band[1])]
    if overlap[1]-overlap[0] <= 1.e-6:
        raise ValueError('COLUMN_HOOP_TIE has no valid axial overlap')
    local = [overlap[0]-start,overlap[1]-start]
    cuts = [s for s in local if 1.e-6 < s < length-1.e-6]
    member['interval'] = local
    return dict(rule_id='COLUMN_HOOP_TIE',enabled=True,group_id=next(iter(groups)),column=member,halves=halves,
        expected_hoop_band=band,expected_overlap=overlap,local_overlap=local,partitions=cuts,tolerance_m=1.e-6)
