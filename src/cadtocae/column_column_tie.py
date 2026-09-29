"""Tie the entire coaxial PIPE overlap, independent of the frozen beam rule."""
import json
import math
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from .analysis_defaults import open_analysis_config
from .analysis_project import validate_project_files
from .workbook import export_abaqus_json
from .column_axis import cct_placement_axis, cct_axial_layout


def read_column_config(path):
    wb = open_analysis_config(path, load_workbook, 'Column_Tie_Config', 'COLUMN_COLUMN_TIE', ('rule_id','enabled'), read_only=True, data_only=False)
    try:
        # Old beam-only workbooks remain beam-only.
        if 'Column_Tie_Config' not in wb.sheetnames:
            return None
        rows = list(wb['Column_Tie_Config'].values)
        if len(rows) != 2 or rows[0] != ('rule_id','enabled'):
            raise ValueError('Column_Tie_Config requires rule_id, enabled (no column_tie_length_mm)')
        rule, enabled = rows[1]
        if rule != 'COLUMN_COLUMN_TIE' or enabled not in (True,False,'TRUE','FALSE') or type(enabled) in (int,float):
            raise ValueError('Column_Tie_Config requires COLUMN_COLUMN_TIE and TRUE/FALSE')
        return dict(rule_id=rule,enabled=enabled in (True,'TRUE'))
    finally:
        wb.close()


def build_column_plan(files, analysis_path):
    config = read_column_config(analysis_path)
    if config is None or not config['enabled']:
        return None
    summary = validate_project_files(files)
    with tempfile.TemporaryDirectory(prefix='step05_column_') as tmp:
        path = export_abaqus_json(files.components, Path(tmp)/'components.json', selection='complete')
        components = json.loads(path.read_text(encoding='utf-8'))['components']
    return resolve_columns(summary, components)


def resolve_columns(summary, components):
    result = dict(rule_id='COLUMN_COLUMN_TIE', enabled=True, region_rule='ENTIRE_ACTUAL_OVERLAP', tolerance_m=1.e-6,
                  tie='STEP05_TIE_COLUMN_UP_DOWN', members=[])
    for role, side, region in (
        ('COLUMN_UP', 'OUTER', 'STEP05_REGION_COLUMN_UP_DOWN'),
        ('COLUMN_DOWN', 'INNER', 'STEP05_REGION_COLUMN_DOWN_UP'),
    ):
        matches = [m for m in summary['instances'] if (m.get('canonical_role') or m.get('component_code')) == role]
        if len(matches) != 1:
            raise ValueError('Expected unique canonical role '+role)
        member = matches[0]
        if len([m for m in summary['instances'] if m['part_name'] == member['part_name']]) != 1:
            raise ValueError('Shared column Part is unsupported')
        matches = [c for c in components if c['part_name'] == member['part_name'] and c['component_code'] == member['component_code']]
        if len(matches) != 1:
            raise ValueError('Missing/ambiguous normalized PIPE component '+role)
        geometry = matches[0]
        if geometry['section_kind'] != 'PIPE' or geometry['model_policy'] != 'SHELL':
            raise ValueError('Expected PIPE shell '+role)
        length = geometry['length_m']
        params = geometry['section_params_m']
        diameter, thickness = params['od_m'], geometry['thickness_m']
        for v in (length, diameter, thickness):
            if isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) or v <= 0:
                raise ValueError('Invalid PIPE dimensions')
        if diameter <= 2*thickness or abs(params['t_m']-thickness) > 1.e-6:
            raise ValueError('Invalid PIPE thickness')
        if abs(member['geometry_info']['length']-length) > 1.e-6:
            raise ValueError('Column workbook/summary length mismatch')
        placement = dict(member['placement'])
        if member.get('axis_direction') is not None:
            placement.setdefault('axis_direction',member['axis_direction'])
        axis = cct_placement_axis(placement)
        origin = placement.get('final_origin',placement['translation'])
        if len(origin) != 3 or any(not isinstance(x, (int,float)) or not math.isfinite(x) for x in origin):
            raise ValueError('Invalid column translation')
        result['members'].append(dict(role=role, part_name=member['part_name'], instance_name=member['instance_name'],
            length_m=length, radius_m=(diameter-thickness)/2., thickness_m=thickness, section_params_m=params,
            placement=placement, origin=list(origin), axis_direction=list(axis),physical_side=side,region=region))
    up, down = result['members']
    result['master'], result['slave'] = down['instance_name'], up['instance_name']
    layout = cct_axial_layout(result['members'],result['tolerance_m'])
    result.update(layout)
    for i, member in enumerate(result['members']):
        member['interval'] = layout['local_overlaps'][i]
        member['partitions'] = layout['partitions'][i]
        member['global_range'] = layout['upper_global_range' if i == 0 else 'lower_global_range']
        member['selected_range'] = list(layout['actual_overlap'])
    return result
