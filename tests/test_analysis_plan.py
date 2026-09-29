from __future__ import annotations

import copy
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook

from cadtocae.analysis_inputs import normalize_inputs, read_analysis_inputs, REGION_COLUMNS
from cadtocae.analysis_plan import semantic_name
from cadtocae.analysis_rules import build_analysis_plan


def benchmark_inputs():
    return {"Project_Info": {"project_id": "SYNTHETIC_A2", "schema_version": 1, "length_unit": "m"},
            "Partition_Region": [{"region_id": "REGION_BEAM_BRACE_FRONT", "connection_id": "BEAM_BRACE_FRONT",
                                  "target_role": "INCLINED_BEAM", "reference_station": "C", "half_length_m": 0.080,
                                  "partition_rule": "BEAM_LOCAL_PATCH", "enabled": True}]}


def benchmark_summary():
    return {"project_id": "SYNTHETIC_A2", "model_name": "SYNTHETIC_A2", "structure_type": "SYNTHETIC",
            "assembly_info": {"executor_version": "step04-instance-plan-v1"},
            "assembly_points": {"C": [0., 0.1, 1.]},
            "instances": [{"instance_name": role+"_01", "part_name": "P_"+role,
                           "canonical_role": role, "component_code": role, "connection_group": None,
                           "geometry_info": {"length": 2., "section_kind": "C_CHANNEL"},
                           "placement": {"translation": [0., 0., 0.], "rotation_steps": []}}
                          for role in ("INCLINED_BEAM", "BRACE_FRONT")]}


class AnalysisInputPlanTests(unittest.TestCase):
    def setUp(self):
        self.inputs, self.summary = benchmark_inputs(), benchmark_summary()

    def test_normalization(self):
        row = self.inputs["Partition_Region"][0]
        row["enabled"], row["half_length_m"] = " ON ", "0.080"
        self.assertEqual(normalize_inputs(self.inputs)["Partition_Region"][0]["half_length_m"], .08)

    def test_missing_station(self):
        self.summary["assembly_points"].clear()
        with self.assertRaisesRegex(ValueError, "station C"):
            build_analysis_plan(self.summary, self.inputs)

    def test_invalid_half_lengths(self):
        for value in (0, -1, float("nan"), float("inf"), True, "wrong"):
            with self.subTest(value=value), self.assertRaises(ValueError):
                data = copy.deepcopy(self.inputs)
                data["Partition_Region"][0]["half_length_m"] = value
                build_analysis_plan(self.summary, data)

    def test_unresolved_and_ambiguous_role(self):
        for instances in (self.summary["instances"][:1], self.summary["instances"] + [dict(self.summary["instances"][1], instance_name="OTHER")]):
            with self.assertRaisesRegex(ValueError, "role"):
                build_analysis_plan(dict(self.summary, instances=instances), self.inputs)

    def test_names_and_disabled_future_objects(self):
        plan = build_analysis_plan(self.summary, self.inputs).to_dict()
        self.assertEqual(semantic_name("REGION_BEAM_BRACE_FRONT"), "STEP05_REGION_BEAM_BRACE_FRONT")
        self.assertEqual(semantic_name("STEP05_REGION_BEAM_BRACE_FRONT"), "STEP05_REGION_BEAM_BRACE_FRONT")
        for key in ("ties", "reference_points", "couplings"):
            self.assertEqual(plan[key], [])
        self.assertEqual(len(plan["partitions"]), 3)
        with self.assertRaises(ValueError):
            semantic_name("Surface-1")

    def test_duplicate_definitions_fail(self):
        self.inputs["Partition_Region"] *= 2
        with self.assertRaisesRegex(ValueError, "Duplicate"):
            build_analysis_plan(self.summary, self.inputs)

    def test_no_structure_type_dispatch(self):
        expected = build_analysis_plan(self.summary, self.inputs).to_dict()
        for kind in ("SP_SC", "SP_DC", "DP"):
            actual = build_analysis_plan(dict(self.summary, structure_type=kind), self.inputs).to_dict()
            self.assertEqual(expected, actual)

    def test_rear_connection_not_enabled(self):
        self.inputs["Partition_Region"][0]["connection_id"] = "BEAM_BRACE_REAR"
        with self.assertRaisesRegex(ValueError, "Unsupported A2"):
            build_analysis_plan(self.summary, self.inputs)

    def test_mismatch_missing_transform_and_units(self):
        with self.assertRaisesRegex(ValueError, "mismatch"):
            build_analysis_plan(dict(self.summary, model_name="OTHER"), self.inputs)
        self.summary["instances"][0]["placement"].pop("rotation_steps")
        with self.assertRaisesRegex(ValueError, "transform contract"):
            build_analysis_plan(self.summary, self.inputs)
        self.inputs["Project_Info"]["length_unit"] = "mm"
        with self.assertRaises(ValueError):
            normalize_inputs(self.inputs)

    def test_xlsx_reader_and_file_closed(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/"analysis.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Project_Info"
            ws.append(["key", "value"])
            for key, value in self.inputs["Project_Info"].items():
                ws.append([key, value])
            ws = wb.create_sheet("Partition_Region")
            ws.append(REGION_COLUMNS)
            ws.append([self.inputs["Partition_Region"][0][key] for key in REGION_COLUMNS])
            wb.save(path)
            self.assertEqual(read_analysis_inputs(path), normalize_inputs(self.inputs))
            ws.cell(2, 5, "=0.080")
            wb.save(path)
            with self.assertRaisesRegex(ValueError, "formulas"):
                read_analysis_inputs(path)
            wb.close()
            path.unlink()  # exercises Windows handle release on both paths
