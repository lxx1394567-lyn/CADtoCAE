"""SP_DC composition and native-geometry contract tests (no Abaqus installed)."""
import copy
import contextlib
import io
import json
from pathlib import Path
from types import SimpleNamespace as NS
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
import test_analysis_beam_brace_tie as bb
import test_analysis_column_hoop_tie as hh
from cadtocae.analysis_defaults import family_analysis_config
from cadtocae.analysis_discovery import scan_projects,project_files
from cadtocae.analysis_batch import generate_batch
from cadtocae.sp_dc_rules import column_beam_plan,hoop_group_plans,normalized_components
from cadtocae.sp_dc_setup import build_sp_dc_plan,generate
from cadtocae.runtime_compatibility import assert_runtime_compatible

ROOT=Path(__file__).resolve().parents[1]

class DCTests(unittest.TestCase):
    def setUp(self):
        self.base=bb.PlannerTests();self.addCleanup(self.base.doCleanups);self.base.setUp()
        self.root=self.base.root;s=self.base.summary;s['structure_type']='SP_DC'
        beam=s['instances'][0]
        beam['placement']=dict(translation=[0,0,.1],axis_direction=[1,0,0],rotation_steps=[])
        hs,hc=hh.inputs('HOOP_ASSEMBLY_1')
        for m in hs['instances']:m.setdefault('canonical_role',m['component_code'])
        for m in s['instances']:m.setdefault('canonical_role',m['component_code'])
        lower=next(m for m in hs['instances'] if m['canonical_role']=='COLUMN_DOWN')
        s['instances'].append(lower)
        halves=[m for m in hs['instances'] if m['canonical_role']=='HOOP']
        s['instances'].extend(halves)
        for h in halves:
            h=copy.deepcopy(h);h.update(component_code='HOOP_ASSEMBLY_2',part_name='P_HOOP2',connection_group='HOOP_02',instance_name=h['instance_name'].replace('01','02'))
            h['placement']['translation'][2]=.9;s['instances'].append(h)
        for role,x in (('COLUMN_FRONT',.4),('COLUMN_REAR',1.3)):
            s['instances'].append(dict(canonical_role=role,component_code=role,part_name='P_'+role,instance_name=role+'_01',
                geometry_info=dict(length=.4,section_kind='C_CHANNEL'),placement=dict(translation=[x,0,-.3],rotation_steps=[])))
            s['assembly_points']['UF1' if role=='COLUMN_FRONT' else 'UR1']=[x,0,.1]
        self.base.write_summary()
        wb=load_workbook(self.base.components)
        for role,label in (('COLUMN_FRONT','前立柱'),('COLUMN_REAR','后立柱')):
            wb.active.append(['单桩双立柱',28,label,'C90x50x15x2',400,1,'Q355B','SHELL','S4R','P_'+role,'已确认'])
        wb.active.append(['单桩双立柱',28,'下立柱','φ177×4',1600,1,'Q355B','SHELL','S4R',lower['part_name'],'已确认'])
        for n,part in ((1,'P_HOOP'),(2,'P_HOOP2')):
            wb.active.append(['单桩双立柱',28,'抱箍组合'+str(n),'HOOP(D=177,W=80,T=6,L=100,R=100,RF=6)',None,1,'Q355B','SOLID','C3D8R',part,'已确认'])
        wb.save(self.base.components);wb.close()
        wb=load_workbook(self.base.coordinate)
        for key in ('UF1','UR1'):wb.active.append([key]+s['assembly_points'][key])
        wb.save(self.base.coordinate);wb.close()
        self.scan=lambda:scan_projects(self.root)
        self.files=project_files(self.scan()['projects'][0])

    def test_capability_without_up_or_f(self):
        r=self.scan()['projects'][0]
        self.assertEqual(r['input_status'],'READY',r['errors']);self.assertEqual(r['capability'],'SP_DC_FULL_CONNECTION_SETUP')
        self.assertEqual(r['expected_tie_count'],8)
        self.assertEqual(r['hoop_groups'],['HOOP_02','HOOP_01'])

    def test_missing_roles(self):
        for role in ('COLUMN_FRONT','COLUMN_REAR','COLUMN_DOWN'):
            original=self.base.summary['instances']
            self.base.summary['instances']=[m for m in original if m['canonical_role']!=role];self.base.write_summary()
            r=self.scan()['projects'][0];self.assertIn(role,r['missing_requirements']['roles']);self.assertEqual(r['input_status'],'NEEDS_REVIEW')
            self.base.summary['instances']=original

    def test_missing_points(self):
        for key in ('C','E','UF1','UR1'):
            p=self.base.summary['assembly_points'].pop(key);self.base.write_summary()
            self.assertIn(key,self.scan()['projects'][0]['missing_requirements']['points'])
            self.base.summary['assembly_points'][key]=p

    def test_plan_eight_ties_no_rp_no_coupling(self):
        p=build_sp_dc_plan(self.files)
        self.assertEqual(p['expected_constraints']['tie_count'],8)
        self.assertEqual(len(set(p['expected_constraints']['ties'])),8)
        self.assertNotIn('column_column_tie',p);self.assertNotIn('column_beam_coupling',p)
        self.assertEqual(p['expected_constraints']['reference_points'],0);self.assertEqual(p['expected_constraints']['couplings'],0)
        self.assertEqual(p['source_files']['analysis_config'],'DEFAULT_SP_DC_ANALYSIS_CONFIG')

    def test_stations_and_partitions(self):
        p=build_sp_dc_plan(self.files);c=p['column_beam_tie']
        for x,m in zip((.4,1.3),c['connections']):
            self.assertAlmostEqual(m['station_m'],x)
            for a,b in zip(m['beam_range'],[x-.05,x+.05]):self.assertAlmostEqual(a,b)
            for a,b in zip(m['column_range'],[.35,.4]):self.assertAlmostEqual(a,b)
            self.assertEqual(m['master'],m['column']);self.assertEqual(m['slave'],m['beam'])
            self.assertAlmostEqual(p['partitions'][m['column_part']][0],.35)

    def test_projected_ang28_numeric_fixture(self):
        s=copy.deepcopy(self.base.summary);components=normalized_components(self.files.components)
        data={m['canonical_role']:m for m in s['instances']}
        beam=data['INCLINED_BEAM'];beam['placement']=dict(translation=[-1.8100425653608,-.045,2.50658329628892],rotation_steps=[dict(axis_direction=[0,1,0],angle_deg=62)])
        beam['geometry_info']['length']=4.062
        for role,x,length in (('COLUMN_FRONT',-.255,1.23),('COLUMN_REAR',.255,1.508)):
            data[role]['placement']['translation']=[x,-.045,2.1];data[role]['geometry_info']['length']=length
        for c in components:
            if c['component_code'] in ('INCLINED_BEAM','COLUMN_FRONT','COLUMN_REAR'):c['length_m']=data[c['component_code']]['geometry_info']['length']
        c=column_beam_plan(s,components,family_analysis_config('SP_DC')['COLUMN_BEAM_TIE'])
        for expected,actual in zip((1.7595918165937368,2.340408183406267),c['connections']):self.assertAlmostEqual(expected,actual['station_m'])

    def test_outside_station_fails(self):
        for m in self.base.summary['instances']:
            if m['canonical_role']=='COLUMN_FRONT':m['placement']['translation'][0]=10
        self.base.write_summary()
        self.assertEqual(self.scan()['projects'][0]['input_status'],'NEEDS_REVIEW')
        with self.assertRaisesRegex(ValueError,'outside Part'):build_sp_dc_plan(self.files)

    def test_incomplete_hoop(self):
        self.base.summary['instances']=[m for m in self.base.summary['instances'] if m['instance_name']!='HOOP_02_B'];self.base.write_summary()
        self.assertTrue(self.scan()['projects'][0]['missing_requirements']['hoop_groups'])

    def test_single_hoop_six_ties(self):
        self.base.summary['instances']=[m for m in self.base.summary['instances'] if m.get('connection_group')!='HOOP_02'];self.base.write_summary()
        self.assertEqual(build_sp_dc_plan(self.files)['expected_constraints']['tie_count'],6)

    def test_overlapping_groups_rejected(self):
        for m in self.base.summary['instances']:
            if m.get('connection_group')=='HOOP_02':m['placement']['translation'][2]=1.27
        self.base.write_summary()
        with self.assertRaisesRegex(ValueError,'Overlapping HOOP'):build_sp_dc_plan(self.files)

    def test_generated_compatibility_and_order(self):
        from lib2to3.refactor import RefactoringTool
        paths=generate(self.files,output_dir=self.root/'out')
        source=paths['script'].read_text(encoding='utf-8');assert_runtime_compatible(source)
        RefactoringTool([]).refactor_string(source+'\n','spdc')
        self.assertNotIn('model.Coupling(',source);self.assertNotIn('.ReferencePoint(',source)
        body=source[source.index('def execute_sp_dc'):]
        self.assertLess(body.index('prepare_hoop('),body.index('execute_column_beam('))
        self.assertLess(body.index('execute_hoop('),body.index('batch.finish('))

    def test_batch_defaults_and_no_xlsx(self):
        before=set(self.root.glob('*.xlsx'))
        r=generate_batch(self.scan())
        self.assertEqual(r['generation_counts'],dict(succeeded=1,failed=0),r['projects'][0]['errors'])
        self.assertEqual(set(self.root.glob('*.xlsx')),before)
        self.assertEqual(self.scan()['projects'][0]['generation_status'],'GENERATED')


class ContactTests(unittest.TestCase):
    def setUp(self):
        self.ns={}
        for name in ('beam_brace_tie_runtime.py','column_beam_tie_runtime.py'):
            source=(ROOT/'src/cadtocae'/name).read_text(encoding='utf-8').split("\nif __name__ == '__main__':")[0]
            exec(source,self.ns)
    def face(self,p,n):return NS(getNormal=lambda:n,getCentroid=lambda:(p,))
    def test_unique_pair_normal_gap_and_center(self):
        beam=NS(faces=[self.face((0,0,0),(0,1,0)),self.face((0,0,0),(1,0,0))])
        column=NS(faces=[self.face((0,.01,0),(0,-1,0)),self.face((0,.02,0),(0,1,0))])
        self.assertEqual(self.ns['cbt_contact_pair'](beam,[0,1],column,[0,1])[2:],(0,0))
    def test_ambiguous_pair_fails(self):
        a=NS(faces=[self.face((0,0,0),(0,1,0))]);b=NS(faces=[self.face((0,.01,0),(0,1,0))]*2)
        with self.assertRaisesRegex(ValueError,'contact face ambiguous'):self.ns['cbt_contact_pair'](a,[0],b,[0,1])
    def test_reversed_normals_change_physical_side(self):
        other=self.face((0,1,0),(0,1,0))
        for normal,side in (((0,1,0),'SIDE1'),((0,-1,0),'SIDE2')):
            self.assertEqual(self.ns['cbt_contact_side'](self.face((0,0,0),normal),other),side)
    def test_coplanar_opposing_sections_resolve(self):
        f=self.face((0,0,0),(0,1,0))
        self.assertEqual(self.ns['cbt_contact_side'](f,f,(0,-1,0),(0,1,0)),'SIDE1')
        with self.assertRaisesRegex(ValueError,'physical side ambiguous'):self.ns['cbt_contact_side'](f,f,(0,1,0),(0,2,0))

class MultiHoopRuntimeTests(unittest.TestCase):
    def test_two_groups_partition_first_four_ties_and_repeat(self):
        fixture=hh.HoopTests();self.addCleanup(fixture.doCleanups);fixture.setUp()
        ns=fixture.runtime
        fixture.base.report.write_text(json.dumps(dict(status='SUCCESS',connections=[])))
        exec((ROOT/'src/cadtocae/tie_preflight_runtime.py').read_text(encoding='utf-8'),ns)
        class MultiPipe(hh.Pipe):
            def __init__(self,member):
                super().__init__(member);self.bands=[]
            def PartitionFaceByDatumPlane(self,datumPlane,faces):
                result=super().PartitionFaceByDatumPlane(datumPlane,faces)
                if not hasattr(datumPlane,'offset'):self.bands.append(self.split)
                return result
            def make_faces(self,origin,instance,axis=None):
                faces=hh.cc.Faces()
                for j,(lo,hi) in enumerate(zip(self.stations,self.stations[1:])):
                    halves=(1,-1) if any(lo>=a-1.e-6 and hi<=b+1.e-6 for a,b in self.bands) else (0,)
                    for half in halves:
                        ids=(2*j,2*j+1,2*j+2,2*j+3) if half else (2*j,2*j+2)
                        faces.append(hh.CurvedFace(len(faces),self.member['radius_m'],lo,hi,half,ids,origin,instance,self.sign))
                return faces
        pipe=MultiPipe(fixture.pipe.member);fixture.model.parts[pipe.name]=pipe
        fixture.a.instances['COLUMN_DOWN_01']=hh.Instance(pipe)
        rule1=copy.deepcopy(fixture.rule);rule2=copy.deepcopy(fixture.rule)
        rule2['group_id']='HOOP_02';rule2['expected_hoop_band']=[.9,.98]
        for rule in (rule1,rule2):
            group=rule['group_id'];rule['column']['partition_namespace']=group
            for half in rule['halves']:
                label=half['label'];half['instance_name']=group+'_'+label
                half['tie']='STEP05_TIE_COLUMN_DOWN_'+group+'_'+label
                half['column_region']='STEP05_REGION_COLUMN_DOWN_'+group+'_'+label
                half['hoop_region']='STEP05_REGION_'+group+'_'+label+'_COLUMN_DOWN'
                if group=='HOOP_02':
                    h=hh.Hoop(label);h.name=half['instance_name']
                    for v in h.vertices:v.pointOn=((v.pointOn[0][0],v.pointOn[0][1],v.pointOn[0][2]-.36),)
                    for face in h.faces:
                        face.instanceName=h.name;face.lo-=.36;face.hi-=.36
                        face.box['low']=tuple(v-(.36 if i==2 else 0) for i,v in enumerate(face.box['low']))
                        face.box['high']=tuple(v-(.36 if i==2 else 0) for i,v in enumerate(face.box['high']))
                        point=face.pointOn[0];face.pointOn=((point[0],point[1],point[2]-.36),face.radial)
                    fixture.a.instances[h.name]=h
        beam=bb.RuntimeTests();self.addCleanup(beam.doCleanups);beam.setUp()
        from analysis_fake_cae import Instance
        class RichInstance(Instance):
            @property
            def faces(self):
                result=[]
                for f in self.part.faces:
                    coords=[self.part.vertices[i].pointOn[0] for i in f.getVertices()]
                    centroid=tuple(sum(p[i] for p in coords)/len(coords) for i in range(3))
                    result.append(NS(index=f.index,instanceName=self.name,getVertices=f.getVertices,
                        getNormal=lambda f=f:self.mapped(f.getNormal(),False),
                        getCentroid=lambda centroid=centroid:(self.mapped(centroid),)))
                return result
        for name,inst in beam.model.rootAssembly.instances.items():
            fixture.a.instances[name]=RichInstance(name,inst.part,inst.rotation,inst.translation)
        fixture.model.parts.update(beam.model.parts)
        fixture.base.mdb.models['TEST']=fixture.model
        plan=copy.deepcopy(beam.plan)
        columns=[]
        identity=((1,0,0),(0,1,0),(0,0,1))
        for suffix,station in (('FRONT',.85),('REAR',1.15)):
            name='COLUMN_'+suffix
            part=type(beam.parts[0])(name,.03,.05,.01,.5)
            fixture.model.parts[name]=part
            fixture.a.instances[name]=RichInstance(name,part,identity,(.001,.02,station-.475))
            plan['instances'].append(dict(part_name=name,instance_name=name,component_geometry=dict(length_m=.5,section_params_m=dict(h_m=.05,b_m=.03))))
            plan['partitions'][name]=[.45]
            plan['partitions']['BEAM'].extend([station-.05,station+.05])
            columns.append(dict(id='BEAM_COLUMN_'+suffix,column=name,beam='BEAM',column_part=name,beam_part='BEAM',
                column_range=[.45,.5],beam_range=[station-.05,station+.05],
                column_region='STEP05_REGION_'+name+'_BEAM',beam_region='STEP05_REGION_BEAM_'+name,
                tie='STEP05_TIE_BEAM_COLUMN_'+suffix))
        plan.update(column_beam_tie=dict(connections=columns),column_hoop_groups=[rule2,rule1],warnings=['fixture warning'])
        for name in ('column_beam_tie_runtime.py','sp_dc_runtime.py'):
            exec((ROOT/'src/cadtocae'/name).read_text(encoding='utf-8'),ns)
        for repeat in (False,True):
            with contextlib.redirect_stdout(io.StringIO()):
                ns['execute_sp_dc'](plan,fixture.base.mdb,str(fixture.base.report),'multi','XYPLANE','ON','COMPUTED')
            report=json.loads(fixture.base.report.read_text())
            self.assertEqual(report['overall_status'],'SUCCESS')
            self.assertEqual(len(fixture.model.constraints),8)
            self.assertEqual(report['warnings'],['fixture warning'])
            self.assertEqual(len(report['column_beam_tie']['connections']),2)
            self.assertEqual(report['slave_preflight']['checked_ties'],8)
            self.assertEqual(len(pipe.bands),2)
            if repeat:self.assertEqual(report['slave_preflight']['created'],[])
            else:self.assertEqual(len(report['slave_preflight']['created']),8)
            previous=report['column_hoop_groups']
            self.assertEqual(len(previous),2)
            self.assertTrue(all(r['status']=='SUCCESS' for r in previous.values()))
            master_ids=[]
            for r in previous.values():
                for c in r['connections']:master_ids+=c['column_face_ids']
            self.assertEqual(len(master_ids),len(set(master_ids)))

if __name__=='__main__':unittest.main()
