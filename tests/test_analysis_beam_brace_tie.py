import ast
import contextlib
import importlib.util
import io
import json
from pathlib import Path
import tempfile
import subprocess
import sys
from types import SimpleNamespace as NS
import unittest
from unittest.mock import patch
from openpyxl import Workbook, load_workbook

from cadtocae.beam_brace_tie import FIELDS, read_tie_config, local_stations, build_plan, generate
from cadtocae.analysis_project import discover_project_folder
import test_analysis_project as project_fixture
from analysis_fake_cae import Part, Instance, Repository


class PlannerTests(unittest.TestCase):
    write_summary = project_fixture.RealProjectTests.write_summary

    def setUp(self):
        project_fixture.RealProjectTests.setUp(self)
        rear = json.loads(json.dumps(self.summary['instances'][1]))
        rear.update(component_code='BRACE_REAR',canonical_role='BRACE_REAR',instance_name='BRACE_REAR_01',part_name='P_BRACE_REAR')
        self.summary['instances'].append(rear)
        self.summary['assembly_points'].update(B=[0.,0.,0.],D=[1.,0.,0.],E=[1.6,0.,.1])
        self.write_summary()
        wb=load_workbook(self.components)
        wb.active.append(['单桩单立柱',18,'后斜撑','C40x50x10x4',1000,1,'Q235B','SHELL','S4R','P_BRACE_REAR','已确认'])
        wb.save(self.components)
        wb.close()
        wb=load_workbook(self.coordinate)
        for name in ('B','D','E'):
            wb.active.append([name]+self.summary['assembly_points'][name])
        s=wb.create_sheet('Parameters')
        s.append(['参数名','数值','单位'])
        s.append(['GC_mm',1000,'mm'])
        s.append(['GE_mm',1600,'mm'])
        wb.save(self.coordinate)
        wb.close()
        self.config=self.root/'analysis.xlsx'
        wb=Workbook()
        wb.active.title='Tie_Config'
        wb.active.append(FIELDS)
        wb.active.append(['BEAM_BRACE_TIE',True,50,50,None,None])
        wb.save(self.config)
        wb.close()

    def test_plan_both_roles_defaults_and_stations(self):
        plan=build_plan(discover_project_folder(self.root),self.config)
        self.assertEqual(len(plan['connections']),2)
        self.assertEqual(plan['config']['beam_shell_side'],'SIDE2')
        for actual, expected in zip(plan['partitions']['P_INCLINED_BEAM'],[.95,1.05,1.55,1.65]):
            self.assertAlmostEqual(actual,expected)
        for c in plan['connections']:
            self.assertEqual(c['master'],c['brace'])
            self.assertEqual(c['slave'],c['beam'])
            self.assertEqual(c['brace_range'],[.95,1.])

    def test_invalid_inputs_and_role_ambiguity(self):
        wb=load_workbook(self.config)
        wb.active['C2']=0
        wb.save(self.config)
        wb.close()
        with self.assertRaises(ValueError): read_tie_config(self.config)
        wb=load_workbook(self.config)
        wb.active['C2']=50
        wb.save(self.config)
        wb.close()
        self.summary['instances'].append(dict(self.summary['instances'][0]))
        self.write_summary()
        with self.assertRaisesRegex(ValueError,'unique canonical role'):
            build_plan(discover_project_folder(self.root),self.config)

    def test_new_cli_default_discovers_analysis(self):
        script=Path(__file__).resolve().parents[1]/'scripts/step05_generate_analysis_setup.py'
        result=subprocess.run([sys.executable,str(script),'--project-folder',str(self.root)],capture_output=True,text=True)
        self.assertEqual(result.returncode,0,result.stderr)
        plan=json.loads(next((self.root/'step05_validation').glob('*_analysis_plan.json')).read_text())
        self.assertEqual(plan['rule_id'],'BEAM_BRACE_TIE')
        self.assertEqual(len(plan['connections']),2)

    def test_config_side_override_and_disabled(self):
        wb=load_workbook(self.config)
        wb.active['E2']='SIDE1'
        wb.active['B2']=False
        wb.save(self.config)
        wb.close()
        plan=build_plan(discover_project_folder(self.root),self.config)
        self.assertEqual(plan['config']['beam_shell_side'],'SIDE1')
        self.assertEqual(plan['connections'],[])

    def test_contact_point_mismatch(self):
        self.summary['assembly_points']['E'][0]+=1
        self.write_summary()
        with self.assertRaisesRegex(ValueError,'point mismatch'):
            build_plan(discover_project_folder(self.root),self.config)

    def test_coordinate_missing_fail_no_pose_fallback(self):
        wb=load_workbook(self.coordinate)
        del wb['Parameters']
        wb.save(self.coordinate)
        wb.close()
        with self.assertRaisesRegex(ValueError,'Missing real GC_mm'):
            local_stations(self.coordinate)

    def test_generated_path_uses_simple_runtime_only(self):
        paths=generate(discover_project_folder(self.root),self.config)
        source=paths['script'].read_text(encoding='utf-8')
        calls={n.func.attr for n in ast.walk(ast.parse(source)) if isinstance(n,ast.Call) and isinstance(n.func,ast.Attribute)}
        self.assertIn('Tie',calls)
        self.assertIn('DatumPlaneByPrincipalPlane',calls)
        self.assertNotIn('PartitionFaceByExtendFace',calls)
        self.assertNotIn('_s5_pose',source)
        self.assertNotIn('placement_candidates',source)
        self.assertFalse(calls & {'Coupling','ReferencePoint','CutExtrude','RemoveFaces'})


class RuntimeTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.report=Path(self.tmp.name)/'report.json'
        path=Path(__file__).resolve().parents[1]/'src/cadtocae/beam_brace_tie_runtime.py'
        self.runtime={'__name__':'test','json':json}
        exec(path.read_text(encoding='utf-8'),self.runtime)
        class TestPart(Part):
            def DatumPlaneByPrincipalPlane(self,principalPlane,offset):
                return self.feature('DatumPlane',NS(point=(0,0,offset),normal=(0,0,1)))
        class TestInstance(Instance):
            @property
            def faces(self):
                return [NS(index=f.index,instanceName=self.name,getVertices=f.getVertices,
                           getNormal=lambda f=f:self.mapped(f.getNormal(),False)) for f in self.part.faces]
        beam=TestPart('BEAM',.05,.09,.015,2.)
        front=TestPart('FRONT',.03,.07,.015,.5)
        rear=TestPart('REAR',.03,.07,.015,.6)
        self.parts=[beam,front,rear]
        identity=((1,0,0),(0,1,0),(0,0,1))
        inst={p.name:TestInstance(p.name,p,identity,(0,0,z)) for p,z in zip(self.parts,(0,.025,.925))}
        a=NS(instances=inst,sets={},surfaces={},regenerate=lambda:None)
        a.Set=lambda name,faces:a.sets.update({name:NS(name=name,faces=faces)})
        def surface(name,**kwargs):
            faces=kwargs.get('side1Faces',kwargs.get('side2Faces'))
            a.surfaces[name]=NS(name=name,faces=faces,sides=['SIDE1' if 'side1Faces' in kwargs else 'SIDE2']*len(faces))
        a.Surface=surface
        model=NS(parts={p.name:p for p in self.parts},rootAssembly=a,constraints={})
        def tie(**kwargs):
            self.assertTrue(kwargs['master'].faces)
            self.assertTrue(kwargs['slave'].faces)
            self.assertEqual(len(a.surfaces)>=2,True)
            model.constraints[kwargs['name']]=NS(**kwargs)
        model.Tie=tie
        self.model=model
        self.mdb=NS(models={'TEST':model})
        self.plan=dict(project_id='TEST',model_name='TEST',config={'enabled':True,'beam_shell_side':'SIDE2','brace_shell_side':'SIDE2'},
                       instances=[dict(part_name=p.name,instance_name=p.name,component_geometry={'length_m':L,'section_params_m':{'h_m':h,'b_m':b}})
                                  for p,L,h,b in zip(self.parts,(2.,.5,.6),(.09,.07,.07),(.05,.03,.03))],
                       partitions={'BEAM':[.45,.55,1.45,1.55],'FRONT':[.05,.45],'REAR':[.05,.55]},connections=[])
        for role,station,L in [('FRONT',.5,.5),('REAR',1.5,.6)]:
            self.plan['connections'].append(dict(id='BEAM_BRACE_'+role,station='C' if role=='FRONT' else 'E',station_m=station,
                beam='BEAM',brace=role,beam_part='BEAM',brace_part=role,beam_range=[station-.05,station+.05],brace_range=[L-.05,L],
                beam_region='STEP05_REGION_BEAM_BRACE_'+role,brace_region='STEP05_REGION_BRACE_'+role+'_BEAM',tie='STEP05_TIE_BEAM_BRACE_'+role))

    def run_script(self):
        with contextlib.redirect_stdout(io.StringIO()):
            return self.runtime['execute'](self.plan,self.mdb,str(self.report),'fixture','XYPLANE','ON','COMPUTED')

    def test_partitions_ties_and_idempotent_reuse(self):
        result=self.run_script()
        self.assertEqual(result['status'],'SUCCESS')
        self.assertEqual(len(result['partitions']),8)
        self.assertEqual(len(self.model.constraints),2)
        for role in ('FRONT','REAR'):
            tie=self.model.constraints['STEP05_TIE_BEAM_BRACE_'+role]
            self.assertIn('BRACE_'+role+'_BEAM',tie.master.name)
            self.assertIn('BEAM_BRACE_'+role,tie.slave.name)
        repeated=self.run_script()
        self.assertEqual(repeated['partitions'],[])
        self.assertFalse(repeated['tie_created'])
        self.assertTrue(repeated['remesh_required'])

    def test_independent_model_rejected_before_modification(self):
        self.model.rootAssembly.instances['FRONT'].dependent='OFF'
        with self.assertRaisesRegex(ValueError,'fresh Step04'):
            self.run_script()
        self.assertFalse(any(p.mutations for p in self.parts))

    def test_empty_or_nonparallel_brace_region_stops_all_ties(self):
        self.plan['connections'][1]['brace_range']=[9,10]
        with self.assertRaisesRegex(ValueError,'Empty'):
            self.run_script()
        self.assertFalse(self.model.constraints)

    def test_conflicting_surface_and_tie_fail(self):
        self.run_script()
        surface=self.model.rootAssembly.surfaces['STEP05_SURF_BEAM_BRACE_FRONT']
        surface.sides=['SIDE1']
        with self.assertRaisesRegex(ValueError,'Conflicting Surface'):
            self.run_script()

    def test_ambiguous_contact_fails(self):
        self.run_script()
        a=self.model.rootAssembly
        fn=self.runtime['contact_faces']
        beamids=self.runtime['zone_faces'](self.parts[0],[.45,.55],True)
        braceids=self.runtime['zone_faces'](self.parts[1],[.45,.5])
        # Repeating the exact nearest face is deliberately ambiguous.
        best=fn(a.instances['BEAM'],beamids,a.instances['FRONT'],braceids)[0][0]
        with self.assertRaisesRegex(ValueError,'Ambiguous'):
            fn(a.instances['BEAM'],beamids,a.instances['FRONT'],[best,best])
