from __future__ import annotations

import ast
import json
import math
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from cadtocae.assembly_script import build_assembly_summary, generate_assembly_scripts_from_workbook
from cadtocae.main_frame_assembly import (
    FAILED, _read_named_points, _dp_hoop_orientation, _create_hoop_pair,
    build_abaqus_script, build_payload, detect_structure_type,
)


PROJECT = "DP_ANG14_TEST"
# Final values from the updated DP input, including the signed -80 mm rear offset.
POINTS = {
    "F": [0.0, 0.0, 2.130], "G": [-1.94059145255199, 0.0, 1.64615620880066],
    "CF0": [-1.4, 0.0, 0.0], "CF1": [-1.4, 0.0, 1.811],
    "CR0": [1.4, 0.0, 0.0], "CR1": [1.4, 0.0, 2.449],
    "HF": [-1.4, 0.0, 0.5], "B": [-1.3, 0.0, 0.5],
    "C": [-0.0194059145255199, 0.0, 2.12516156208801],
    "HR": [1.4, 0.0, 0.53], "D": [1.32, 0.0, 0.53],
    "E": [0.0776236581020797, 0.0, 2.14935375164797],
    "P1": [-1.79504709361059, 0.0, 1.68244449314061],
    "P2": [-0.436633076824198, 0.0, 2.02113514698015],
    "P3": [0.434692485371646, 0.0, 2.23838100922865],
    "P4": [1.79310650215804, 0.0, 2.57707166306819],
}


def component(code, length, kind, params):
    return {"component_code": code, "part_name": "P_TEST_" + code, "length_m": length,
            "section_kind": kind, "section_params_m": params}


class DPAdapterTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.excel = self.root / (PROJECT + "_coordinate.xlsx")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "坐标计算总表"
        for name, value in (("theta_deg", 14.0), ("length_tolerance_mm", 10.0),
                            ("angle_tolerance_deg", 0.1), ("L_front_brace_mm", 2080.0),
                            ("L_rear_brace_mm", 2080.0)):
            sheet.append([name, "input", value])
        sheet.append(["关键控制点输出"])
        sheet.append(["点名", "X_m", "Y_m", "Z_m", "状态"])
        for name, xyz in POINTS.items():
            sheet.append([name, *xyz, "已就绪"])
        sheet.append([])
        sheet.append(["校验结果（长度、角度）"])
        sheet.append(["校核项", "计算值", "图纸/输入值", "误差", "允许误差", "是否通过"])
        for name, actual, reference in (("FRONT_BRACE_LENGTH", 2.06907499009595, 2.08),
                                        ("REAR_BRACE_LENGTH", 2.04103051125749, 2.08),
                                        ("COLUMN_TOP_LINE_ANGLE", 12.8361019635766, 14.0)):
            sheet.append([name, actual, reference, actual - reference, 0.01, "不通过"])
        workbook.save(self.excel)
        workbook.close()
        pipe = {"od_m": 0.06, "t_m": 0.002}
        channel = {"h_m": 0.08, "b_m": 0.04, "lip_m": 0.015, "t_m": 0.002}
        self.components = [
            component("COLUMN_FRONT", 1.811, "PIPE", pipe),
            component("COLUMN_REAR", 2.509, "PIPE", dict(pipe)),
            component("INCLINED_BEAM", 4.0, "C_CHANNEL", channel),
            component("BRACE_FRONT", 2.08, "C_CHANNEL", dict(channel)),
            component("BRACE_REAR", 2.08, "C_CHANNEL", dict(channel)),
            component("HOOP", None, "HOOP_BAND", {"diameter_m": 0.06, "width_m": 0.04, "t_m": 0.003,
                                                     "left_extension_m": 0.057, "right_extension_m": 0.117}),
            component("PURLIN", 16.326, "C_CHANNEL", channel),
            component("PURLIN_SUPPORT", 0.05, "ANGLE", {"leg_a_m": 0.09, "leg_b_m": 0.063, "t_m": 0.006}),
            component("PURLIN_LOCAL", 0.05, "C_CHANNEL", channel),
        ]
        self.part_script = self.root / (PROJECT + "_create_parts_in_cae.py")

    def payload(self):
        self.part_script.write_text(
            "MODEL_NAME = %r\nCOMPONENTS_JSON = r'''%s'''\n" % (PROJECT, json.dumps({"components": self.components})),
            encoding="utf-8",
        )
        return build_payload(self.excel, self.part_script)

    def update_point(self, name, xyz):
        workbook = load_workbook(self.excel)
        sheet = workbook.active
        for row in sheet:
            if row[0].value == name:
                for col, value in enumerate(xyz, start=2):
                    sheet.cell(row[0].row, col, value)
                break
        workbook.save(self.excel)
        workbook.close()

    def assertVector(self, actual, expected):
        for a, b in zip(actual, expected):
            self.assertAlmostEqual(a, b, places=9)

    def test_dispatch_main_frame_with_four_purlin_pairs(self):
        payload = self.payload()
        self.assertEqual(detect_structure_type(PROJECT), "DP")
        self.assertEqual(payload["meta"]["structure_type"], "DP")
        self.assertEqual([m["instance_id"] for m in payload["instance_plan"][:9]], [
            "COLUMN_FRONT_01", "COLUMN_REAR_01", "INCLINED_BEAM_01", "BRACE_FRONT_01", "BRACE_REAR_01",
            "HOOP_01_A", "HOOP_01_B", "HOOP_02_A", "HOOP_02_B",
        ])
        self.assertEqual(len(payload["required_part_names"]), 8)
        self.assertEqual(len(payload["purlin_nodes"]), 4)
        self.assertEqual(len(payload["instance_plan"]), 17)
        for code in ("COLUMN_DOWN", "PURLIN"):
            self.assertNotIn(code, [m["component_code"] for m in payload["instance_plan"]])
            self.assertNotIn("P_TEST_" + code, payload["required_part_names"])

    def test_columns_and_braces_use_exact_final_control_lines_without_scaling(self):
        plan = {m["component_code"]: m for m in self.payload()["instance_plan"]}
        for code, start, end, length in (("COLUMN_FRONT", "CF0", "CF1", 1.811),
                                        ("COLUMN_REAR", "CR0", "CR1", 2.509),
                                        ("BRACE_FRONT", "B", "C", 2.08), ("BRACE_REAR", "D", "E", 2.08)):
            with self.subTest(code=code):
                member = plan[code]
                self.assertEqual(member["placement_mode"], "LINE")
                self.assertEqual(member["origin"], POINTS[start])
                self.assertEqual(member["start_point"], POINTS[start])
                self.assertEqual(member["end_point"], POINTS[end])
                self.assertVector(member["axis_direction"], [(b-a)/math.dist(POINTS[start], POINTS[end])
                                                            for a, b in zip(POINTS[start], POINTS[end])])
                self.assertEqual(member["part_length_m"], length)
        self.assertEqual(plan["BRACE_REAR"]["origin"][0], 1.32)

    def test_generated_common_transforms_keep_anchors_axes_and_actual_lengths(self):
        payload = self.payload()
        script = build_abaqus_script("full_main_frame", payload, Path("report.json"), Path("unused.cae"))
        compile(script, "generated_dp.py", "exec")
        tree = ast.parse(script)
        functions = ast.Module(body=[n for n in tree.body if isinstance(n, ast.FunctionDef)], type_ignores=[])
        namespace = {"math": math}
        exec(compile(functions, "generated_helpers", "exec"), namespace)
        transform = namespace["_transform_member"]
        for member in payload["instance_plan"][:5]:
            anchor = member["local_anchor"]
            actual_start = transform(anchor, member)
            actual_end = transform([anchor[0], anchor[1], member["part_length_m"]], member)
            self.assertVector(actual_start, member["origin"])
            self.assertAlmostEqual(math.dist(actual_start, actual_end), member["part_length_m"])
            self.assertVector([(b-a)/member["part_length_m"] for a,b in zip(actual_start,actual_end)], member["axis_direction"])
            if member["component_code"] == "COLUMN_REAR":
                self.assertAlmostEqual(actual_end[2] - POINTS["CR1"][2], 0.060)

    def test_beam_uses_g_f_and_preserves_roll(self):
        beam = next(m for m in self.payload()["instance_plan"] if m["component_code"] == "INCLINED_BEAM")
        self.assertEqual(beam["origin"], POINTS["G"])
        self.assertAlmostEqual(beam["rotate_y_deg"], 76.0)
        self.assertEqual(beam["roll_about_axis_deg"], 90.0)
        self.assertEqual(beam["part_length_m"], 4.0)

    def test_hoop_pairs_reuse_part_with_horizontal_flip_through_hf_hr(self):
        payload = self.payload()
        for index, point, column in ((1, "HF", "COLUMN_FRONT_01"), (2, "HR", "COLUMN_REAR_01")):
            pair = [m for m in payload["instance_plan"] if m.get("hoop_group_id") == "HOOP_%02d" % index]
            self.assertEqual(len(pair), 2)
            for member in pair:
                self.assertEqual(member["part_name"], "P_TEST_HOOP")
                self.assertEqual(member["origin"], POINTS[point])
                self.assertEqual(member["column_instance"], column)
                self.assertEqual(member["control_point"], point)
            self.assertEqual(pair[1]["post_rotation"], {
                "axis_point": POINTS[point], "axis_direction": payload["hoop_groups"][index-1]["long_plate_direction"], "angle_deg": 180.0,
            })
            self.assertAlmostEqual(pair[1]["post_rotation"]["axis_direction"][2], 0.0)

    def generated_helpers(self, payload):
        script = build_abaqus_script("full_main_frame", payload, Path("report.json"), Path("unused.cae"))
        tree = ast.parse(script)
        functions = ast.Module(body=[n for n in tree.body if isinstance(n, ast.FunctionDef)], type_ignores=[])
        namespace = {"math": math}
        exec(compile(functions, "generated_helpers", "exec"), namespace)
        return namespace

    def test_actual_transforms_align_centers_long_short_plates_arcs_and_width(self):
        payload = self.payload()
        helpers = self.generated_helpers(payload)
        transform = helpers["_transform_member"]
        vector = helpers["_transform_member_vector"]
        for group in payload["hoop_groups"]:
            with self.subTest(group=group["group_id"]):
                a, b = [m for m in payload["instance_plan"] if m.get("hoop_group_id") == group["group_id"]]
                inward = [1.0 if group["column_role"] == "FRONT" else -1.0, 0.0, 0.0]
                for member in (a, b):
                    self.assertVector(transform([0, 0, 0.02], member), group["center"])
                    self.assertVector(vector([1, 0, 0], member), inward)
                    self.assertVector(vector([-1, 0, 0], member), [-x for x in inward])
                    limits = sorted(transform([0, 0, z], member)[2] for z in (0.0, 0.04))
                    self.assertVector(limits, [group["center"][2] - 0.02, group["center"][2] + 0.02])
                self.assertVector(vector([0, 1, 0], a), [-x for x in vector([0, 1, 0], b)])
                # Matching outer plate edges: B reverses the width coordinate, not X.
                for x in (-0.090, 0.150):
                    for z in (0.0, 0.02, 0.04):
                        self.assertVector(transform([x, 0, z], a), transform([x, 0, 0.04-z], b))
                self.assertEqual(group["orientation_check"], "PASS")
                self.assertAlmostEqual(group["brace_side_dot_product"], 0.1 if group["column_role"] == "FRONT" else 0.08)

    def test_column_derived_direction_and_flip_work_away_from_global_x(self):
        # Exercise the small orientation rule independently of the current XZ-only frame adapter.
        for own, other, expected in (([1, 2, 9], [4, 6, -7], [0.6, 0.8, 0.0]),
                                     ([4, 6, -7], [1, 2, 9], [-0.6, -0.8, 0.0])):
            orientation, inward, _ = _dp_hoop_orientation(self.components[5], own, other)
            self.assertVector(inward, expected)
            pair, _ = _create_hoop_pair(self.components[5], [1, 2, 3], 1, "test",
                                        initial_orientation=orientation, pair_axis_direction=inward)
            helpers = self.generated_helpers(self.payload())
            for member in pair:
                self.assertVector(helpers["_transform_member_vector"]([1, 0, 0], member), expected)
                self.assertVector(helpers["_transform_member"]([0, 0, 0.02], member), [1, 2, 3])

    def test_front_rear_directions_follow_swapped_column_positions(self):
        for name, point in POINTS.items():
            self.update_point(name, [-point[0], point[1], point[2]])
        groups = self.payload()["hoop_groups"]
        self.assertVector(groups[0]["long_plate_direction"], [-1, 0, 0])
        self.assertVector(groups[1]["long_plate_direction"], [1, 0, 0])
        self.assertTrue(all(g["orientation_check"] == "PASS" for g in groups))

    def test_wrong_or_zero_brace_side_warns_and_exports_diagnostics(self):
        for name, xyz, index in (("B", [-1.5, 0, 0.5], 0), ("D", [1.4, 0, 0.53], 1)):
            self.update_point(name, xyz)
            payload = self.payload()
            group = build_assembly_summary(payload)["hoop_groups"][index]
            self.assertEqual(group["orientation_check"], "WARNING")
            self.assertLessEqual(group["brace_side_dot_product"], 0.0)
            self.assertTrue(any(group["group_id"] + "_ORIENTATION" in w for w in payload["warnings"]))

    def test_invalid_dp_geometry_or_coincident_columns_fails_preflight(self):
        with self.assertRaisesRegex(ValueError, "horizontal positions coincide"):
            _dp_hoop_orientation(self.components[5], [0, 0, 0], [0, 0, 9])
        self.components[5]["section_params_m"]["right_extension_m"] = 0.0
        with self.assertRaisesRegex(ValueError, "finite positive"):
            self.payload()

    def test_symmetric_or_nearly_symmetric_dp_hoop_is_ambiguous(self):
        for delta in (0.0, 0.5e-6, -0.5e-6):
            with self.subTest(delta=delta):
                self.components[5]["section_params_m"]["right_extension_m"] = 0.057 + delta
                with self.assertRaisesRegex(ValueError, "asymmetric long-plate direction cannot be resolved"):
                    self.payload()

    def test_ang14_and_ang35_actual_long_side_drives_both_half_transforms(self):
        for left, right, width, sign in ((0.057, 0.117, 0.04, 1), (0.08, 0.05, 0.05, -1)):
            with self.subTest(left=left, right=right):
                params = self.components[5]["section_params_m"]
                params.update(left_extension_m=left, right_extension_m=right, width_m=width)
                payload = self.payload()
                helpers = self.generated_helpers(payload)
                point = helpers["_transform_member"]
                vector = helpers["_transform_member_vector"]
                for index, group in enumerate(payload["hoop_groups"]):
                    local_long = [sign, 0, 0]
                    self.assertEqual(group["local_long_plate_direction"], local_long)
                    self.assertEqual(group["local_short_plate_direction"], [-sign, 0, 0])
                    self.assertEqual(group["left_extension_m"], left)
                    self.assertEqual(group["right_extension_m"], right)
                    expected = [1 if index == 0 else -1, 0, 0]
                    expected_yaw = (0 if index == 0 else 180) if sign == 1 else (180 if index == 0 else 0)
                    self.assertEqual(group["initial_orientation"], [{"axis": "Z", "angle_deg": expected_yaw}])
                    self.assertVector(group["long_plate_direction_global"], expected)
                    self.assertVector(group["pair_axis_direction"], expected)
                    self.assertGreater(group["brace_side_dot_product"], 0)
                    a, b = [m for m in payload["instance_plan"] if m.get("hoop_group_id") == group["group_id"]]
                    for member in (a, b):
                        self.assertVector(vector(local_long, member), expected)
                        self.assertVector(vector([-sign, 0, 0], member), [-v for v in expected])
                        self.assertVector(point([0, 0, width/2], member), group["center"])
                        limits = sorted(point([0, 0, z], member)[2] for z in (0, width))
                        self.assertVector(limits, [group["center"][2]-width/2, group["center"][2]+width/2])
                    self.assertVector(vector([0, 1, 0], a), [-v for v in vector([0, 1, 0], b)])
                    for x in (-left-0.033, right+0.033):
                        self.assertVector(point([x, 0, 0], a), point([x, 0, width], b))

    def test_left_long_hoop_aligns_to_non_x_inward_direction(self):
        self.components[5]["section_params_m"].update(left_extension_m=0.08, right_extension_m=0.05)
        orientation, inward, geometry = _dp_hoop_orientation(self.components[5], [1, 2, 9], [4, 6, -7])
        pair, _ = _create_hoop_pair(self.components[5], [1, 2, 3], 1, "test",
                                    initial_orientation=orientation, pair_axis_direction=inward)
        helpers = self.generated_helpers(self.payload())
        for member in pair:
            self.assertVector(helpers["_transform_member_vector"](geometry["local_long_plate_direction"], member), [0.6, 0.8, 0])

    def test_qa_and_rear_column_mismatch_are_warnings_only(self):
        payload = self.payload()
        self.assertEqual(payload["errors"], [])
        self.assertEqual(len(payload["warnings"]), 4)
        for key, error in (("FRONT_BRACE_LENGTH", -0.01092500990405),
                           ("REAR_BRACE_LENGTH", -0.03896948874251), ("COLUMN_REAR_LENGTH", -0.060)):
            self.assertAlmostEqual(payload["checks"][key]["error_m"], error)
            self.assertEqual(payload["checks"][key]["passed"], FAILED)
            self.assertTrue(any(key in warning for warning in payload["warnings"]))
        self.assertAlmostEqual(payload["checks"]["COLUMN_TOP_LINE_ANGLE"]["error_deg"], -1.16389803642341)

    def test_named_reader_stops_before_numeric_qa_rows(self):
        points, _ = _read_named_points(self.excel)
        self.assertEqual(set(points), set(POINTS))

    def test_purlin_nodes_use_saved_points_real_parts_and_physical_surfaces(self):
        # Different ANGLE dimensions and lengths must not require angle-specific placement.
        for leg_a, leg_b, thickness, radius in ((0.09, 0.063, 0.006, 0.007), (0.075, 0.05, 0.005, 0.006)):
            with self.subTest(leg_a=leg_a):
                self.components[7].update(part_name="ACTUAL_SUPPORT_PART", length_m=0.07)
                self.components[7]["section_params_m"] = {"leg_a_m": leg_a, "leg_b_m": leg_b, "t_m": thickness, "root_radius_m": radius}
                self.components[8].update(part_name="ACTUAL_LOCAL_PART", length_m=0.06)
                payload = self.payload()
                transform = self.generated_helpers(payload)["_transform_member"]
                summary = build_assembly_summary(payload)
                self.assertEqual(len(summary["purlin_nodes"]), 4)
                for index, node in enumerate(payload["purlin_nodes"], 1):
                    pair = payload["instance_plan"][9+2*(index-1):11+2*(index-1)]
                    support, local = pair
                    self.assertEqual([m["instance_id"] for m in pair], ["PURLIN_SUPPORT_%02d" % index, "PURLIN_LOCAL_%02d" % index])
                    self.assertEqual([m["part_name"] for m in pair], ["ACTUAL_SUPPORT_PART", "ACTUAL_LOCAL_PART"])
                    self.assertTrue(all(m["connection_group"] == "P%02d" % index for m in pair))
                    station = POINTS["P%d" % index]
                    self.assertEqual(node["station_xyz"], station)
                    n, tangent = node["beam_normal"], node["beam_tangent"]
                    beam_top = [station[i] + 0.041*n[i] + (0.02 if i == 1 else 0) for i in range(3)]
                    support_surface = transform([0, 0, 0.035], support)
                    local_midplane = transform([0, 0, 0.03], local)
                    self.assertVector(support_surface, beam_top)
                    self.assertAlmostEqual(sum((local_midplane[i]-beam_top[i])*n[i] for i in range(3))-0.002/2, 0.010)
                    # Both web and support outer long-leg plane have zero tangent separation.
                    self.assertAlmostEqual(sum((local_midplane[i]-support_surface[i])*tangent[i] for i in range(3)), 0.0)
                    self.assertAlmostEqual(support_surface[1], station[1]+0.02)
                    self.assertAlmostEqual(local_midplane[1], station[1]+0.02)
                    self.assertEqual(node["web_contact_side"], "outer")
                    self.assertAlmostEqual(node["alignment_error_m"], 0)
                self.assertNotEqual(payload["checks"]["PURLIN_STATION_ORDER"]["passed"], FAILED)

    def test_nonzero_transverse_beam_position_translates_purlin_nodes(self):
        original = self.payload()
        for name, xyz in POINTS.items():
            self.update_point(name, [xyz[0], xyz[1]+0.37, xyz[2]])
        shifted = self.payload()
        for a, b in zip(original["purlin_nodes"], shifted["purlin_nodes"]):
            for key in ("beam_section_center", "support_length_midpoint", "local_purlin_length_midpoint"):
                self.assertVector(b[key], [a[key][0], a[key][1]+0.37, a[key][2]])
            self.assertAlmostEqual(b["alignment_error_m"], 0)

    def test_reversed_saved_stations_warn_without_repositioning(self):
        self.update_point("P1", POINTS["P2"])
        self.update_point("P2", POINTS["P1"])
        payload = self.payload()
        self.assertEqual(payload["checks"]["PURLIN_STATION_ORDER"]["passed"], FAILED)
        self.assertEqual(payload["purlin_nodes"][0]["station_xyz"], POINTS["P2"])
        self.assertTrue(any("PURLIN_STATION_ORDER" in w for w in payload["warnings"]))

    def test_summary_preserves_points_warnings_and_hoop_ownership(self):
        summary = build_assembly_summary(self.payload())
        self.assertEqual(summary["structure_type"], "DP")
        self.assertEqual(len(summary["instances"]), 17)
        for name, point in POINTS.items():
            self.assertEqual(summary["assembly_points"][name], point)
        self.assertFalse(set(summary["checks"]) & set(summary["assembly_points"]))
        self.assertEqual(len(summary["warnings"]), 4)
        self.assertEqual([(g["control_point"], g["column_instance"]) for g in summary["hoop_groups"]],
                         [("HF", "COLUMN_FRONT_01"), ("HR", "COLUMN_REAR_01")])
        self.assertEqual(summary["instances"][8]["connection_group"], "HOOP_02")
        self.assertEqual(summary["instances"][-1]["connection_group"], "P04")

    def test_real_generation_entry_exports_despite_warnings(self):
        self.payload()
        result = generate_assembly_scripts_from_workbook(self.excel, self.root / "out", components_json=self.part_script)
        self.assertEqual(result.status, "needs_review")
        self.assertEqual(result.error_count, 0)
        self.assertEqual(result.warning_count, 4)
        self.assertEqual(len(result.script_paths), 1)
        self.assertIsNone(result.assembly_json_path)
        summary = json.loads(Path(result.summary_path).read_text(encoding="utf-8"))
        report = json.loads(Path(result.report_path).read_text(encoding="utf-8"))
        self.assertEqual(summary["warnings"], report["warnings"])
        self.assertEqual(summary["hoop_groups"], report["hoop_groups"])
        required = {"group_id", "column_role", "center", "part_name", "inward_direction", "long_plate_direction",
                    "short_plate_direction", "pair_transform_rule", "brace_role", "brace_side_dot_product", "orientation_check"}
        self.assertTrue(all(required <= set(group) for group in summary["hoop_groups"]))
        self.assertEqual(summary["assembly_points"]["D"], [1.32, 0.0, 0.53])

    def test_numbered_hoops_match_by_diameter_not_index(self):
        self.components[1]["section_params_m"]["od_m"] = 0.08
        self.components = [c for c in self.components if c["component_code"] != "HOOP"]
        self.components.extend([
            component("HOOP_ASSEMBLY_2", None, "HOOP_BAND", {"diameter_m": 0.06, "width_m": 0.04, "left_extension_m": 0.057, "right_extension_m": 0.117}),
            component("HOOP_ASSEMBLY_1", None, "HOOP_BAND", {"diameter_m": 0.08, "width_m": 0.04, "left_extension_m": 0.057, "right_extension_m": 0.117}),
        ])
        groups = self.payload()["hoop_groups"]
        self.assertEqual([g["source_component_code"] for g in groups], ["HOOP_ASSEMBLY_2", "HOOP_ASSEMBLY_1"])

    def test_ambiguous_hoop_mapping_is_preflight_error(self):
        self.components.append(component("HOOP_ASSEMBLY_8", None, "HOOP_BAND", {"diameter_m": 0.06}))
        with self.assertRaisesRegex(ValueError, "ambiguous"):
            self.payload()

    def test_missing_matching_hoop_is_preflight_error(self):
        self.components[5]["section_params_m"]["diameter_m"] = 0.09
        with self.assertRaisesRegex(ValueError, "mapping missing"):
            self.payload()

    def test_control_points_are_authoritative_not_rebuilt_from_symmetry(self):
        self.update_point("CF0", [-1.6, 0.0, 0.2])
        self.update_point("CF1", [-1.6, 0.0, 2.011])
        self.update_point("HF", [-1.6, 0.0, 0.7])
        self.update_point("D", [1.29, 0.0, 0.6])
        payload = self.payload()
        self.assertEqual(payload["instance_plan"][0]["origin"], [-1.6, 0.0, 0.2])
        self.assertEqual(payload["instance_plan"][4]["origin"], [1.29, 0.0, 0.6])
        self.assertEqual(payload["instance_plan"][2]["origin"], POINTS["G"])

    def test_nonvertical_column_or_off_axis_hoop_is_rejected(self):
        self.update_point("CF1", [-1.3, 0.0, 1.811])
        with self.assertRaisesRegex(ValueError, "vertical"):
            self.payload()
        self.update_point("CF1", POINTS["CF1"])
        self.update_point("HF", [-1.3, 0.0, 0.5])
        with self.assertRaisesRegex(ValueError, "own COLUMN_FRONT axis"):
            self.payload()


if __name__ == "__main__":
    unittest.main()
