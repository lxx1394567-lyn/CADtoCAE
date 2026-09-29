"""Real-mode workflow tests use temporary synthetic fixtures, never real files."""
import ast
import contextlib
import hashlib
import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from cadtocae.analysis_project import (discover_project_folder, explicit_project_files,
                                     build_real_analysis_plan, generate_real_analysis_setup)
from analysis_fake_cae import fake_model_and_summary

PROJECT = "SP_SC_ANG18_TEST_FIXTURE"


class RealProjectTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.mdb, self.summary = fake_model_and_summary()
        self.mdb.models[PROJECT] = self.mdb.models.pop("SYNTHETIC_A2")
        self.summary.update(project_id=PROJECT, model_name=PROJECT, structure_type="SP_SC")
        self.summary["assembly_points"]["G_global"] = [0., 0., .1]
        self.summary_path = self.root/(PROJECT+"_assembly_summary.json")
        self.write_summary()
        self.components = self.root/(PROJECT+"_components.xlsx")
        wb = Workbook()
        sheet = wb.active
        sheet.title = "建模构件表"
        sheet.append(["支架类型", "角度", "构件名称", "规格", "长度_mm", "数量", "材料牌号",
                      "建模方式", "单元类型", "abaqus_part_name", "校核状态"])
        sheet.append(["单桩单立柱", 18, "斜梁", "C200x100x25x4", 2000, 1, "Q235B", "SHELL", "S4R", "P_INCLINED_BEAM", "已确认"])
        sheet.append(["单桩单立柱", 18, "前斜撑", "C40x50x10x4", 1000, 1, "Q235B", "SHELL", "S4R", "P_BRACE_FRONT", "已确认"])
        wb.save(self.components)
        wb.close()
        self.coordinate = self.root/(PROJECT+"_coordinate.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.append(["点名", "X_m", "Y_m", "Z_m"])
        ws.append(["C", 1., 0., .1])
        ws.append(["G_global", 0., 0., .1])
        wb.save(self.coordinate)
        wb.close()
        # Deliberately not Python; reading these to reconstruct data must fail.
        for suffix in ("_create_parts_in_cae.py", "_assembly_frame.py"):
            (self.root/(PROJECT+suffix)).write_bytes(b"NOT PYTHON\x00do not parse")

    def write_summary(self):
        self.summary_path.write_text(json.dumps(self.summary), encoding="utf-8")

    def test_discovery_and_geometry_extraction(self):
        files = discover_project_folder(self.root)
        plan = build_real_analysis_plan(files)
        self.assertEqual(files.components, self.components)
        self.assertEqual(plan["project"]["mode"], "REAL_CASE")
        self.assertEqual(plan["project"]["structure_type"], "SP_SC")
        self.assertEqual(plan["real_case"]["C_global"], [1., 0., .1])
        self.assertEqual(plan["real_case"]["C_beam_local_station_m"], 1.)
        self.assertEqual(plan["real_case"]["patch_range_m"], [.92, 1.08])
        beam = plan["instances"][0]["component_geometry"]
        self.assertEqual(beam["section_params_m"], {"h_m": .2, "b_m": .1, "lip_m": .025, "t_m": .004})
        self.assertEqual(beam["length_m"], 2.)
        self.assertNotIn("SYNTHETIC_A2", json.dumps(plan))

    def test_explicit_mode(self):
        files = explicit_project_files(self.components, self.coordinate, self.summary_path)
        self.assertEqual(files, discover_project_folder(self.root))

    def test_coordinate_status_header_variant(self):
        # The newer coordinate layout has a parameter table and uses 状态
        # on its numeric point table; the legacy all-table reader rejects it.
        wb = load_workbook(self.coordinate)
        ws = wb.active
        ws.cell(1, 5, "状态")
        params = wb.create_sheet("Parameters")
        params.append(["参数名", "参数含义", "数值", "单位", "校核状态", "备注"])
        params.append(["GC_mm", "Station", 1000, "mm", "已确认", ""])
        wb.save(self.coordinate)
        wb.close()
        plan = build_real_analysis_plan(discover_project_folder(self.root))
        self.assertEqual(plan["real_case"]["C_beam_local_station_m"], 1.)
        self.assertIn("verified G_global", plan["real_case"]["source"])

    def test_multiple_candidates_listed(self):
        for filename in ("OTHER_components.xlsx", "OTHER_coordinate.xlsx", "OTHER_assembly_summary.json"):
            path = self.root/filename
            path.write_text("duplicate")
            with self.subTest(filename=filename), self.assertRaisesRegex(ValueError, "ambiguous project files") as raised:
                discover_project_folder(self.root)
            self.assertIn(filename, str(raised.exception))
            path.unlink()

    def test_prefix_and_model_mismatch(self):
        other = self.components.with_name("OTHER_components.xlsx")
        self.components.rename(other)
        with self.assertRaisesRegex(ValueError, "prefix mismatch"):
            discover_project_folder(self.root)
        other.rename(self.components)
        self.summary["model_name"] = "OTHER"
        self.write_summary()
        with self.assertRaisesRegex(ValueError, "model_name mismatch"):
            discover_project_folder(self.root)

    def test_missing_inputs_fail(self):
        files = discover_project_folder(self.root)
        for path in files.paths().values():
            backup = path.with_suffix(path.suffix+".bak")
            path.rename(backup)
            with self.subTest(path=path.name), self.assertRaisesRegex(ValueError, "missing real input"):
                discover_project_folder(self.root)
            backup.rename(path)

    def test_missing_c_no_fallback(self):
        self.summary["assembly_points"].pop("C")
        self.write_summary()
        with self.assertRaisesRegex(ValueError, "station C"):
            generate_real_analysis_setup(discover_project_folder(self.root))
        self.assertFalse((self.root/"step05_validation").exists())

    def test_missing_geometry_no_fallback(self):
        wb = load_workbook(self.components)
        wb.active.cell(2, 4, "C200x100x25")  # no thickness
        wb.save(self.components)
        wb.close()
        with self.assertRaisesRegex(ValueError, "component"):
            generate_real_analysis_setup(discover_project_folder(self.root))
        self.assertFalse((self.root/"step05_validation").exists())

    def test_coordinate_mismatch_rejected(self):
        wb = load_workbook(self.coordinate)
        wb.active.cell(2, 2, 1.01)
        wb.save(self.coordinate)
        wb.close()
        with self.assertRaisesRegex(ValueError, "C mismatch"):
            build_real_analysis_plan(discover_project_folder(self.root))

    def test_component_length_mismatch_rejected(self):
        self.summary["instances"][0]["geometry_info"]["length"] = 2.5
        self.write_summary()
        with self.assertRaisesRegex(ValueError, "length mismatch"):
            build_real_analysis_plan(discover_project_folder(self.root))

    def test_generation_preserves_inputs_and_never_reads_python(self):
        files = discover_project_folder(self.root)
        before = {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in files.paths().values()}
        original = Path.open
        def guarded(path, mode="r", *args, **kwargs):
            if path in (files.part_script, files.assembly_script) and "r" in mode:
                raise AssertionError("Must not read input executable script")
            return original(path, mode, *args, **kwargs)
        with patch.object(Path, "open", guarded):
            results = generate_real_analysis_setup(files)
        self.assertEqual(results["script"].parent, self.root/"step05_validation")
        self.assertEqual(results["script"].name, PROJECT+"_analysis_setup.py")
        source = results["script"].read_text(encoding="utf-8")
        self.assertNotIn("SYNTHETIC_A2", source)
        for name in ("STEP05_REGION_BEAM_BRACE_FRONT", "STEP05_REGION_BRACE_FRONT_BEAM", "STEP05_C_MINUS_80", "STEP05_C_PLUS_80"):
            self.assertIn(name, source)
        calls = {node.func.attr for node in ast.walk(ast.parse(source)) if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute)}
        self.assertFalse(calls & {"Tie", "Coupling", "ReferencePoint", "CutExtrude", "RemoveFace", "ConnectorSection", "Model", "Instance"})
        self.assertFalse(results["runtime_report"].exists())
        self.assertEqual(before, {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in files.paths().values()})

    def test_real_runtime_preflight_and_diagnostics(self):
        results = generate_real_analysis_setup(discover_project_folder(self.root))
        runtime = {"__name__": "test_runtime"}
        exec(compile(results["script"].read_text(encoding="utf-8"), "real_script.py", "exec"), runtime)
        stream = io.StringIO()
        with contextlib.redirect_stdout(stream):
            result = runtime["execute_analysis"](runtime["PLAN"], self.mdb, str(results["runtime_report"]), runtime["PLAN_FINGERPRINT"])
        self.assertEqual(result["status"], "SUCCESS")
        self.assertIn("STEP05 REAL CASE", stream.getvalue())
        self.assertIn("C local station", stream.getvalue())
        self.assertIn("Material removed = False", stream.getvalue())
        names = {s["name"] for s in result["sets"]}
        self.assertTrue({"STEP05_C_MINUS_80", "STEP05_C_PLUS_80"}.issubset(names))

    def test_real_runtime_thickness_mismatch_before_mutation(self):
        results = generate_real_analysis_setup(discover_project_folder(self.root))
        runtime = {"__name__": "test_runtime"}
        exec(compile(results["script"].read_text(encoding="utf-8"), "real_script.py", "exec"), runtime)
        self.mdb.models[PROJECT].sections["shell"].thickness = .005
        with contextlib.redirect_stdout(io.StringIO()), self.assertRaisesRegex(ValueError, "thickness mismatch"):
            runtime["execute_analysis"](runtime["PLAN"], self.mdb, str(results["runtime_report"]), runtime["PLAN_FINGERPRINT"])
        self.assertTrue(all(not p.mutations for p in self.mdb.models[PROJECT].parts.values()))

    def test_cli_folder_and_explicit(self):
        script = Path(__file__).resolve().parents[1]/"scripts/step05_generate_analysis_setup.py"
        for args in (["--project-folder", str(self.root)],
                     ["--components", str(self.components), "--coordinate", str(self.coordinate), "--assembly-summary", str(self.summary_path)]):
            result = subprocess.run([sys.executable, str(script), '--legacy-a2', *args], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("In the same Abaqus/CAE model", result.stdout)

    def test_coordinate_variant_and_lockfile(self):
        (self.root/("~$"+self.components.name)).write_text("Excel lock")
        variant = self.coordinate.with_name(PROJECT+"_coordinate_formula_simple_fixed.xlsx")
        self.coordinate.rename(variant)
        self.assertEqual(discover_project_folder(self.root).coordinate, variant)
