from __future__ import annotations
import json
import shutil
import tempfile
import unittest
import importlib.util
import queue
from types import SimpleNamespace
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
from cadtocae.assembly_projects import scan_projects, generate_batch, can_generate
from cadtocae.assembly_script import generate_assembly_scripts_from_workbook
import test_step04_dp_adapter as dp_fixture
from test_step04_instance_plan import _build_dc_inputs, PROJECT_ID
from test_main_frame_assembly import build_sample_inputs


class DiscoveryTest(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.builder = patch('cadtocae.assembly_projects.build_payload', return_value={
            'instance_plan': [{}], 'warnings': [], 'errors': []}).start()
        self.addCleanup(patch.stopall)

    def pair(self, project='DP_ANG14_TEST', folder=None, metadata=None):
        folder = folder or self.root
        folder.mkdir(parents=True, exist_ok=True)
        part = folder / (project + '_create_parts_in_cae.py')
        data = {'components': [{'component_code': 'HOOP'}], **(metadata or {})}
        part.write_text("MODEL_NAME = %r\nCOMPONENTS_JSON = r'''%s'''\n" % (project, json.dumps(data)), encoding='utf-8')
        coordinate = folder / (project + '_coordinate.xlsx')
        wb = Workbook(); wb.active.append(['project_id', project]); wb.save(coordinate); wb.close()
        return part, coordinate

    def test_valid_types_and_multiple_projects(self):
        for project in ('SP_SC_ANG18_TEST', 'SP_DC_ANG17_TEST', 'DP_ANG35_TEST'):
            self.pair(project)
        rows = scan_projects(self.root)['projects']
        self.assertEqual(len(rows), 3)
        self.assertEqual({r['structure_type'] for r in rows}, {'SP_SC', 'SP_DC', 'DP'})
        self.assertTrue(all(can_generate(r) for r in rows))

    def test_missing_inputs_are_specific(self):
        a, b = self.pair('DP_A'); b.unlink()
        a, b = self.pair('DP_B'); a.unlink()
        self.assertEqual([r['input_status'] for r in scan_projects(self.root)['projects']], ['Missing Coordinate', 'Missing Part Script'])

    def test_recursion_and_duplicate_inputs(self):
        part, coordinate = self.pair()
        nested = self.root/'nested'; nested.mkdir()
        self.pair('SP_DC_CHILD', nested)
        self.assertEqual(len(scan_projects(self.root)['projects']), 1)
        self.assertEqual(len(scan_projects(self.root, True)['projects']), 2)
        for path in (part, coordinate):
            with self.subTest(role=path.suffix):
                shutil.copy2(path, nested/path.name)
                row = scan_projects(self.root, True)['projects'][0]
                self.assertEqual(row['input_status'], 'Ambiguous Files')
                self.assertIn(str(nested/path.name), row['errors'][0])
                (nested/path.name).unlink()

    def test_generated_outputs_and_excel_lock_files_are_ignored(self):
        self.pair()
        for name in ('DP_OTHER_assembly_frame.py', 'DP_OTHER_assembly_summary.json', 'DP_OTHER_step04_assembly_script_report.json', 'DP_OTHER_full_main_frame_report.json', '~$DP_OTHER_coordinate.xlsx', 'test.cae', 'test.exe'):
            (self.root/name).write_text('ignored')
        self.assertEqual(len(scan_projects(self.root)['projects']), 1)

    def test_id_and_structure_metadata_mismatches(self):
        for metadata in ({'project_id': 'DP_OTHER'}, {'meta': {'structure_type': 'SP_SC'}}):
            self.pair(metadata=metadata)
            self.assertEqual(scan_projects(self.root)['projects'][0]['input_status'], 'Project ID Mismatch')
        part, coordinate = self.pair()
        part.write_text(part.read_text().replace("MODEL_NAME = 'DP_ANG14_TEST'", "MODEL_NAME = 'DP_OTHER'"))
        self.assertEqual(scan_projects(self.root)['projects'][0]['input_status'], 'Project ID Mismatch')
        self.pair()
        wb=load_workbook(coordinate); wb.active.cell(1,2,'DP_OTHER'); wb.save(coordinate); wb.close()
        self.assertEqual(scan_projects(self.root)['projects'][0]['input_status'], 'Project ID Mismatch')

    def test_invalid_inputs_and_unsupported_structure(self):
        part, coordinate = self.pair()
        part.write_text('invalid')
        self.assertEqual(scan_projects(self.root)['projects'][0]['input_status'], 'Invalid Part Metadata')
        self.pair(); coordinate.write_text('broken xlsx')
        self.assertEqual(scan_projects(self.root)['projects'][0]['input_status'], 'Invalid Coordinate')
        self.pair('OTHER_A')
        self.assertEqual(scan_projects(self.root)['projects'][1]['capability'], 'Unsupported')

    def test_readonly_scan_warnings_remain_generatable(self):
        self.pair(); self.builder.return_value['warnings'] = ['length mismatch']
        before = sorted(self.root.iterdir())
        row = scan_projects(self.root)['projects'][0]
        self.assertTrue(can_generate(row)); self.assertEqual(row['warnings'], ['length mismatch'])
        self.assertEqual(before, sorted(self.root.iterdir()))

    def test_split_folder_pair_is_not_silently_merged(self):
        part, coordinate = self.pair()
        nested=self.root/'nested'; nested.mkdir(); part.rename(nested/part.name)
        row=scan_projects(self.root, True)['projects'][0]
        self.assertEqual(row['input_status'], 'Ambiguous Files')

    def test_batch_rechecks_added_duplicate_and_continues(self):
        part, coordinate = self.pair('DP_A'); self.pair('DP_B')
        scan=scan_projects(self.root, True)
        nested=self.root/'nested'; nested.mkdir(); shutil.copy2(coordinate,nested/coordinate.name)
        events=[]
        with patch('cadtocae.assembly_projects.generate_assembly_scripts_from_workbook', side_effect=RuntimeError('injected generator failure')) as generate:
            result=generate_batch(scan, progress=events.append)
        self.assertEqual(generate.call_count,1)
        self.assertEqual(result['generation_counts']['failed'],2)
        self.assertEqual([e['row']['project_id'] for e in events if e['event']=='started'],['DP_A','DP_B'])
        self.assertTrue(all(Path(r['report_path']).is_file() for r in result['projects']))

    def test_selected_only_and_all_ready_skips_blocked(self):
        part, coordinate=self.pair('DP_A'); self.pair('DP_B'); coordinate.unlink()
        scan=scan_projects(self.root)
        with patch('cadtocae.assembly_projects.generate_assembly_scripts_from_workbook', side_effect=RuntimeError('test')) as generate:
            result=generate_batch(scan, {'DP_A'})
            generate.assert_not_called()
            self.assertEqual(result['generation_counts']['total'],1)
            result=generate_batch(scan)
            self.assertEqual(generate.call_count,1)
            self.assertEqual(result['generation_counts']['skipped'],1)


class FolderGenerationIntegrationTest(unittest.TestCase):
    def test_gui_selected_dispatch_uses_only_selected_project(self):
        path=Path(__file__).resolve().parents[1]/'scripts/step04_assembly_script_gui.py'
        spec=importlib.util.spec_from_file_location('step04_folder_gui_test',path)
        gui=importlib.util.module_from_spec(spec); spec.loader.exec_module(gui)
        scan={'projects':[{'project_id':'DP_A'},{'project_id':'SP_DC_B'}]}
        app=SimpleNamespace(scan=scan,busy=False,table=SimpleNamespace(selection=lambda:('1',)),
                            status=SimpleNamespace(set=lambda value:None),events=queue.Queue(),
                            run=lambda operation,action:action())
        with patch.object(gui,'generate_batch',return_value=scan) as batch:
            gui.Step04AssemblyScriptApp.start_generate(app,True)
            self.assertEqual(batch.call_args.args[1],{'SP_DC_B'})
            gui.Step04AssemblyScriptApp.start_generate(app,False)
            self.assertIsNone(batch.call_args.args[1])

    def test_three_real_adapter_fixtures_match_direct_file_api(self):
        for structure in ('SP_SC','SP_DC','DP'):
            with self.subTest(structure=structure), tempfile.TemporaryDirectory() as tmp:
                root=Path(tmp)
                if structure=='DP':
                    fixture=dp_fixture.DPAdapterTest(); fixture.setUp()
                    try:
                        fixture.payload()
                        coordinate=root/fixture.excel.name; part=root/fixture.part_script.name
                        shutil.copy2(fixture.excel,coordinate); shutil.copy2(fixture.part_script,part)
                    finally: fixture.doCleanups()
                else:
                    if structure=='SP_SC':
                        excel, components=build_sample_inputs(root); project='SP_SC_ANG20'
                        coordinate=root/(project+'_coordinate.xlsx'); excel.rename(coordinate)
                    else:
                        coordinate, components=_build_dc_inputs(root); project=PROJECT_ID
                    part=root/(project+'_create_parts_in_cae.py')
                    part.write_text("MODEL_NAME = %r\nCOMPONENTS_JSON = r'''%s'''\n" % (project,components.read_text(encoding='utf-8')),encoding='utf-8')
                scan=scan_projects(root)
                self.assertEqual(len(scan['projects']),1)
                self.assertTrue(can_generate(scan['projects'][0]),scan)
                direct=generate_assembly_scripts_from_workbook(coordinate,root/'direct',components_json=part)
                result=generate_batch(scan)
                row=result['projects'][0]
                self.assertEqual(row['generation_status'],'Generated',row)
                expected=json.loads(Path(direct.summary_path).read_text(encoding='utf-8'))
                actual=json.loads(Path(row['generated_files'][1]).read_text(encoding='utf-8'))
                for data in (expected,actual):data['assembly_info'].pop('created_time')
                self.assertEqual(expected,actual)
                for path in row['generated_files']+[row['report_path']]:self.assertEqual(Path(path).parent,root)


if __name__=='__main__':unittest.main()
